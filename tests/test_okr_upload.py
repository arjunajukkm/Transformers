"""
test_okr_upload.py
──────────────────
Tests for the OKR Upload generation mode in KRA Management.
Validates hierarchical Objective-KR association, weightage normalization,
30-column KEKA output format, single-sheet fallback, employee mapping,
blank row handling, header whitespace tolerance, validation errors,
and regression safety of the existing KRA generator.
"""

import os
import tempfile
import pandas as pd
import pytest
from openpyxl import load_workbook

from generate_upload import (
    generate_okr_upload_file,
    generate_upload_file,
    normalize_okr_weightage,
    _find_okr_columns,
    _parse_okr_dataframe,
)


# 1. Hierarchy test
def test_okr_hierarchy():
    df = pd.DataFrame({
        "OKR": ["Objective A", None, None],
        "Key Result Areas": ["Key Result 1", "Key Result 2", "Key Result 3"],
        "Weightage": [0.50, None, None]
    })
    
    extracted = _parse_okr_dataframe(df)
    assert len(extracted) == 3
    for obj_title, kr_title, weight in extracted:
        assert obj_title == "Objective A"
        assert weight == 50
    assert extracted[0][1] == "Key Result 1"
    assert extracted[1][1] == "Key Result 2"
    assert extracted[2][1] == "Key Result 3"


# 2. Multiple Objectives test
def test_okr_multiple_objectives():
    df = pd.DataFrame({
        "OKR": ["Objective A", None, "Objective B", None, "Objective C"],
        "Key Result Areas": ["KR 1", "KR 2", "KR 1", "KR 2", "KR 1"],
        "Weightage": [0.50, None, 0.25, None, 0.25]
    })
    
    extracted = _parse_okr_dataframe(df)
    assert len(extracted) == 5
    assert extracted[0] == ("Objective A", "KR 1", 50)
    assert extracted[1] == ("Objective A", "KR 2", 50)
    assert extracted[2] == ("Objective B", "KR 1", 25)
    assert extracted[3] == ("Objective B", "KR 2", 25)
    assert extracted[4] == ("Objective C", "KR 1", 25)


# 3. Objective without Weightage must NOT inherit previous objective's weight
def test_okr_weightage_no_leak():
    df = pd.DataFrame({
        "OKR": ["Objective A", None, "Objective B", None],
        "Key Result Areas": ["KR 1", "KR 2", "KR 3", "KR 4"],
        "Weightage": [0.50, None, None, None]
    })
    
    extracted = _parse_okr_dataframe(df)
    assert len(extracted) == 4
    assert extracted[0] == ("Objective A", "KR 1", 50)
    assert extracted[1] == ("Objective A", "KR 2", 50)
    assert extracted[2] == ("Objective B", "KR 3", None)
    assert extracted[3] == ("Objective B", "KR 4", None)


# 4. Decimal normalization tests
@pytest.mark.parametrize("input_val, expected", [
    (0.50, 50),
    (0.5, 50),
    (0.25, 25),
    (0.10, 10),
    (1.00, 100),
    (1, 100),
    (50, 50),
    (25, 25),
    (100, 100),
    (50.0, 50),
    (25.0, 25),
    ("0.50", 50),
    ("50%", 50),
    ("0.25", 25),
    (33.33, 33.33),
    (None, None),
])
def test_decimal_normalization(input_val, expected):
    result = normalize_okr_weightage(input_val)
    assert result == expected
    if isinstance(expected, int):
        assert isinstance(result, int)


# 5. Revised headers test (exactly 30 columns, A1:AD1 merge)
def test_revised_headers_and_output_structure(tmp_path):
    input_file = tmp_path / "okr_source.xlsx"
    output_file = tmp_path / "okr_output.xlsx"
    
    df = pd.DataFrame({
        "OKR": ["Objective A", None],
        "Key Result Areas": ["KR 1", "KR 2"],
        "Weightage": [0.50, None]
    })
    df.to_excel(input_file, sheet_name="SSE", index=False)
    
    employees = {
        "SSE": [("EMP001", "John Doe")]
    }
    
    generate_okr_upload_file(str(input_file), str(output_file), employees)
    
    wb = load_workbook(output_file)
    ws = wb["Objective Key Results"]
    
    # Check title row merge
    merge_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert "A1:AD1" in merge_ranges
    assert ws["A1"].value == "Objective Key Results"
    
    # Check header row (row 2)
    headers = [ws.cell(row=2, column=col).value for col in range(1, 35) if ws.cell(row=2, column=col).value is not None]
    assert len(headers) == 30
    assert headers[0] == "Employee Number"
    assert headers[1] == "Employee Name"
    assert headers[2] == "Objective Title"
    assert headers[17] == "Objective Weightage"
    assert headers[18] == "KeyResult Employee Number"
    assert headers[19] == "KeyResult Employee Name"
    assert headers[20] == "Key Result Title"
    assert headers[29] == "Key Result Tags"
    
    # Assert Key Result Weightage is completely absent
    assert "Key Result Weightage" not in headers
    
    # Check data rows (rows 3 and 4)
    assert ws.cell(row=3, column=1).value == "EMP001"
    assert ws.cell(row=3, column=2).value == "John Doe"
    assert ws.cell(row=3, column=3).value == "Objective A"
    assert ws.cell(row=3, column=18).value == 50
    assert ws.cell(row=3, column=21).value == "KR 1"
    
    assert ws.cell(row=4, column=1).value == "EMP001"
    assert ws.cell(row=4, column=2).value == "John Doe"
    assert ws.cell(row=4, column=3).value == "Objective A"
    assert ws.cell(row=4, column=18).value == 50
    assert ws.cell(row=4, column=21).value == "KR 2"
    
    # Row 5 should be empty
    assert ws.cell(row=5, column=1).value is None


# 6. Multiple employees test with alternating row fills
def test_multiple_employees_alternating_rows(tmp_path):
    input_file = tmp_path / "okr_source.xlsx"
    output_file = tmp_path / "okr_output.xlsx"
    
    df = pd.DataFrame({
        "OKR": ["Objective A"],
        "Key Result Areas": ["KR 1"],
        "Weightage": [1.0]
    })
    df.to_excel(input_file, sheet_name="Engineers", index=False)
    
    employees = {
        "Engineers": [
            ("EMP001", "Alice"),
            ("EMP002", "Bob")
        ]
    }
    
    generate_okr_upload_file(str(input_file), str(output_file), employees)
    
    wb = load_workbook(output_file)
    ws = wb["Objective Key Results"]
    
    # Alice on row 3, Bob on row 4
    assert ws.cell(row=3, column=1).value == "EMP001"
    assert ws.cell(row=3, column=2).value == "Alice"
    assert ws.cell(row=4, column=1).value == "EMP002"
    assert ws.cell(row=4, column=2).value == "Bob"
    
    fill_alice = ws.cell(row=3, column=1).fill.start_color.rgb
    fill_bob = ws.cell(row=4, column=1).fill.start_color.rgb
    # One is blue (DCE6F1), the other is white (FFFFFF)
    assert fill_alice != fill_bob


# 7. Blank rows handling
def test_blank_rows_ignored():
    df = pd.DataFrame({
        "OKR": ["Objective A", None, None, None, "Objective B", None],
        "Key Result Areas": ["KR 1", None, "", "KR 2", None, "KR 3"],
        "Weightage": [0.50, None, None, None, 0.25, None]
    })
    
    extracted = _parse_okr_dataframe(df)
    # Only rows with valid non-blank KR should be emitted:
    # "Objective A" with "KR 1"
    # "Objective A" with "KR 2"
    # "Objective B" with "KR 3"
    assert len(extracted) == 3
    assert extracted[0] == ("Objective A", "KR 1", 50)
    assert extracted[1] == ("Objective A", "KR 2", 50)
    assert extracted[2] == ("Objective B", "KR 3", 25)


# 8. Header whitespace tolerance
def test_header_whitespace_tolerance():
    df = pd.DataFrame({
        "  OKR  ": ["Objective A"],
        " Key Result Areas ": ["KR 1"],
        "Weightage   ": [0.50]
    })
    
    okr_col, kr_col, weight_col = _find_okr_columns(df)
    assert okr_col == "  OKR  "
    assert kr_col == " Key Result Areas "
    assert weight_col == "Weightage   "
    
    extracted = _parse_okr_dataframe(df)
    assert len(extracted) == 1
    assert extracted[0] == ("Objective A", "KR 1", 50)


# 9. Single worksheet fallback
def test_single_worksheet_fallback(tmp_path):
    input_file = tmp_path / "single_sheet.xlsx"
    output_file = tmp_path / "single_output.xlsx"
    
    df = pd.DataFrame({
        "OKR": ["Core Objective"],
        "Key Result Areas": ["System Reliability 99.9%"],
        "Weightage": [1.0]
    })
    # Sheet name is generic "Sheet1"
    df.to_excel(input_file, sheet_name="Sheet1", index=False)
    
    # Employee mapping has designation "Senior Software Engineer" which differs from "Sheet1"
    employees = {
        "Senior Software Engineer": [("EMP999", "Dev Hero")]
    }
    
    generate_okr_upload_file(str(input_file), str(output_file), employees)
    
    wb = load_workbook(output_file)
    ws = wb["Objective Key Results"]
    assert ws.cell(row=3, column=1).value == "EMP999"
    assert ws.cell(row=3, column=2).value == "Dev Hero"
    assert ws.cell(row=3, column=3).value == "Core Objective"
    assert ws.cell(row=3, column=18).value == 100
    assert ws.cell(row=3, column=21).value == "System Reliability 99.9%"


# 10. Validation errors
def test_validation_missing_columns():
    df = pd.DataFrame({
        "OKR": ["Objective A"],
        "Key Result Areas": ["KR 1"]
        # Missing Weightage
    })
    with pytest.raises(ValueError, match="missing the required column: Weightage"):
        _parse_okr_dataframe(df)


def test_validation_no_valid_krs():
    df = pd.DataFrame({
        "OKR": ["Objective A"],
        "Key Result Areas": [None],
        "Weightage": [0.50]
    })
    with pytest.raises(ValueError, match="No valid Key Result Areas"):
        _parse_okr_dataframe(df)


def test_validation_file_not_found():
    with pytest.raises(FileNotFoundError):
        generate_okr_upload_file("non_existent_file.xlsx", "out.xlsx", {"Sheet1": [("E1", "A")]})


# 11. Regression safety: original generate_upload_file still works with 31 columns
def test_existing_kra_generator_regression(tmp_path):
    input_file = tmp_path / "kra_source.xlsx"
    output_file = tmp_path / "kra_output.xlsx"
    
    df = pd.DataFrame({
        "KRA": ["Performance", None],
        "KPI": ["Task completion", "Code review"],
        "KRA Weightage": [40, None],
        "KPI Weightage": [20, 20]
    })
    df.to_excel(input_file, sheet_name="Manager", index=False)
    
    employees = {
        "Manager": [("MGR001", "Boss")]
    }
    
    generate_upload_file(str(input_file), str(output_file), employees)
    
    wb = load_workbook(output_file)
    ws = wb["Objective Key Results"]
    
    merge_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert "A1:AE1" in merge_ranges
    
    headers = [ws.cell(row=2, column=col).value for col in range(1, 35) if ws.cell(row=2, column=col).value is not None]
    assert len(headers) == 31
    assert headers[30] == "Key Result Weightage"


# 12. Synonyms: Objective/Goal and Key Results (KRs)
def test_okr_synonyms_objective_goal_and_krs():
    df = pd.DataFrame({
        "Objective/Goal": ["Own Sprint Planning", None],
        "Key Results (KRs)": ["Maintain 95% velocity", "Reduce spillover < 5%"],
        "Weightage": [40.0, None]
    })
    extracted = _parse_okr_dataframe(df)
    assert len(extracted) == 2
    assert extracted[0] == ("Own Sprint Planning", "Maintain 95% velocity", 40)
    assert extracted[1] == ("Own Sprint Planning", "Reduce spillover < 5%", 40)


# 13. Header row offset: Titles/notes on rows 0-2, actual headers on row 3
def test_okr_offset_header_row():
    raw_data = [
        ["Quarterly OKR Draft", None, None],
        ["Owner: Engineering Lead", None, None],
        ["Confidential internal document", None, None],
        ["Objective", "Key Result", "Weight"],
        ["Platform Resiliency", "Achieve 99.95% uptime", 0.30],
        [None, "Zero critical vulnerabilities", None],
        ["Cost Optimization", "Reduce cloud spend by 15%", 0.70]
    ]
    df = pd.DataFrame(raw_data)
    extracted = _parse_okr_dataframe(df)
    assert len(extracted) == 3
    assert extracted[0] == ("Platform Resiliency", "Achieve 99.95% uptime", 30)
    assert extracted[1] == ("Platform Resiliency", "Zero critical vulnerabilities", 30)
    assert extracted[2] == ("Cost Optimization", "Reduce cloud spend by 15%", 70)

