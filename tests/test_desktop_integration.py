"""
tests/test_desktop_integration.py
───────────────────────────────────
Desktop integration tests for Transformers 2.0 — Phase 2B.

Verifies:
 1. Successful upload and snapshot activation.
 2. Dashboard metrics matching the existing analytical engine.
 3. Correct Business Unit, Department, Manager and Employee breakdowns.
 4. Search filtering prepared rows without analytical recalculation.
 5. Repeated reporting-level selection reusing prepared results.
 6. Correct dataset replacement and UI state update.
 7. Failed upload preserving the existing dataset and UI state.
 8. Older processing jobs not overwriting newer completed results (stale job fencing).
 9. No Tkinter widget updates from background worker threads.
10. Excel export continuing to produce identical results and workbook structure.
"""

from datetime import date
from pathlib import Path
import threading
from unittest.mock import MagicMock, patch

import openpyxl
import pandas as pd
import pytest

from app import App
from storage import snapshot_service, TimeSeriesSnapshotService
import time_series_analysis as tsa


@pytest.fixture(scope="module")
def synthetic_excel(tmp_path_factory):
    """Generates a multi-entity synthetic dataset with distinct metrics (9 rows, 3 emps)."""
    tmp_dir = tmp_path_factory.mktemp("integration_data")
    rows = [
        # Emp 1: Alice (Engineering, Backend, RM: Bob Manager) - 3 rows
        {
            "Employee Number": "E001",
            "Employee Name": "Alice",
            "Business Unit": "Engineering",
            "Department": "Backend",
            "Reporting Manager": "Bob Manager",
            "Date": "2026-09-01",
            "Month": "Sep 2026",
            "Attendance Type": "Present",
            "Status": "P",
            "In Time": "09:00",
            "Out Time": "18:00",
            "Applied By": "NA",
            "Applied On": "NA",
            "Approved By": "NA",
            "Approved On": "NA",
        },
        {
            "Employee Number": "E001",
            "Employee Name": "Alice",
            "Business Unit": "Engineering",
            "Department": "Backend",
            "Reporting Manager": "Bob Manager",
            "Date": "2026-09-02",
            "Month": "Sep 2026",
            "Attendance Type": "Missing Swipes",
            "Status": "P(MS)",
            "In Time": "10:00",
            "Out Time": "18:30",
            "Applied By": "NA",
            "Applied On": "NA",
            "Approved By": "NA",
            "Approved On": "NA",
        },
        {
            "Employee Number": "E001",
            "Employee Name": "Alice",
            "Business Unit": "Engineering",
            "Department": "Backend",
            "Reporting Manager": "Bob Manager",
            "Date": "2026-09-03",
            "Month": "Sep 2026",
            "Attendance Type": "Leave",
            "Status": "CL",
            "Leave Name": "Casual Leave",
            "In Time": "NA",
            "Out Time": "NA",
            "Applied By": "Employee",
            "Applied On": "2026-09-01",  # 2 days in advance
            "Approved By": "Manager",
            "Approved On": "2026-09-02",  # 1 day turnaround
        },
        # Emp 2: Charlie (Sales, Direct Sales, RM: David RM) - 1 row
        {
            "Employee Number": "E002",
            "Employee Name": "Charlie",
            "Business Unit": "Sales",
            "Department": "Direct Sales",
            "Reporting Manager": "David RM",
            "Date": "2026-09-01",
            "Month": "Sep 2026",
            "Attendance Type": "Regularized",
            "Status": "AR",
            "In Time": "09:30",
            "Out Time": "18:30",
            "Applied By": "Employee",
            "Applied On": "2026-09-01",
            "Approved By": "Manager",
            "Approved On": "2026-09-02",
        },
        # Emp 2: Charlie WFH 4 days (>3 days threshold) - 4 rows
        *[
            {
                "Employee Number": "E002",
                "Employee Name": "Charlie",
                "Business Unit": "Sales",
                "Department": "Direct Sales",
                "Reporting Manager": "David RM",
                "Date": f"2026-09-0{d}",
                "Month": "Sep 2026",
                "Attendance Type": "Work From Home",
                "Status": "WFH",
                "In Time": "NA",
                "Out Time": "NA",
                "Applied By": "Employee",
                "Applied On": f"2026-09-0{d}",
                "Approved By": "Manager",
                "Approved On": f"2026-09-0{d}",
            }
            for d in range(2, 6)
        ],
        # Emp 3: Eve (Marketing, Growth, RM: Frank RM) - 1 row
        {
            "Employee Number": "E003",
            "Employee Name": "Eve",
            "Business Unit": "Marketing",
            "Department": "Growth",
            "Reporting Manager": "Frank RM",
            "Date": "2026-09-01",
            "Month": "Sep 2026",
            "Attendance Type": "Present",
            "Status": "P",
            "In Time": "08:30",
            "Out Time": "17:30",
            "Applied By": "NA",
            "Applied On": "NA",
            "Approved By": "NA",
            "Approved On": "NA",
        },
    ]
    df = pd.DataFrame(rows)
    p = tmp_dir / "synthetic_dataset.xlsx"
    df.to_excel(p, index=False)
    return p


@pytest.fixture(scope="module")
def replacement_excel(tmp_path_factory):
    """Independent second dataset to test dataset replacement."""
    tmp_dir = tmp_path_factory.mktemp("replacement_data")
    rows = [
        {
            "Employee Number": "E999",
            "Employee Name": "Zack",
            "Business Unit": "Executive",
            "Department": "Board",
            "Reporting Manager": "Chairperson",
            "Date": "2026-11-01",
            "Month": "Nov 2026",
            "Attendance Type": "Present",
            "Status": "P",
            "In Time": "10:00",
            "Out Time": "16:00",
            "Applied By": "NA",
            "Applied On": "NA",
            "Approved By": "NA",
            "Approved On": "NA",
        }
    ]
    df = pd.DataFrame(rows)
    p = tmp_dir / "replacement_dataset.xlsx"
    df.to_excel(p, index=False)
    return p


@pytest.fixture(scope="module")
def desktop_app(synthetic_excel):
    """Headless Desktop App instance reused across integration tests to protect Tcl interpreter."""
    snapshot_service.clear()
    app = App()
    app.withdraw()  # Headless mode

    # Initial load of synthetic dataset
    snap = snapshot_service.prepare_dataset(str(synthetic_excel))
    app._ts_load_success(snap, str(synthetic_excel), job_id=1)
    app.update()

    yield app

    try:
        app.destroy()
    except Exception:
        pass
    snapshot_service.clear()


# =========================================================================
# 1. Successful upload and snapshot activation
# =========================================================================
def test_successful_upload_and_snapshot_activation(desktop_app, synthetic_excel):
    """Verify that dataset ingestion prepares the snapshot and activates UI cleanly."""
    app = desktop_app

    assert snapshot_service.has_active_snapshot()
    assert app.is_ts_loading is False
    assert app.btn_ts_load.cget("state") == "normal"
    assert app.ts_dataset is not None
    assert len(app.ts_dataset) == 9
    assert "Active: 9 records • 3 employees" in app.ts_lbl.cget("text")

    # Check filter dropdown values populated
    bu_vals = app.ts_bu_combo.cget("values")
    assert "All Business Units" in bu_vals
    assert "Engineering" in bu_vals
    assert "Sales" in bu_vals
    assert "Marketing" in bu_vals


# =========================================================================
# 2. Dashboard metrics matching the existing analytical engine
# =========================================================================
def test_dashboard_metrics_matching_existing_analytical_engine(desktop_app):
    """Verify all KPI cards display values matching the ground-truth calculation engine."""
    app = desktop_app
    app._reset_ts_filters()
    app.update()

    # Ground truth
    gt = tsa.compute_time_series_metrics(app.ts_dataset)

    # Check KPI labels in UI
    leave_days = gt['avg_leave_apply_days']
    if leave_days >= 0:
        assert app.lbl_kpi_leave_val.cget("text") == f"{leave_days:.1f} days"
    else:
        assert app.lbl_kpi_leave_val.cget("text") == f"{abs(leave_days):.1f} days late"

    assert app.lbl_kpi_appr_val.cget("text") == f"{gt['avg_approval_days']:.1f} days"
    assert app.lbl_kpi_wfh_val.cget("text") == f"{gt['wfh_excess_days']:,} days"
    assert app.lbl_kpi_rep_val.cget("text") == f"{gt['repeat_emp_count']:,} Emps ({gt['repeat_emp_pct']}%)"
    assert app.lbl_kpi_in_val.cget("text") == gt["avg_in_time"]
    assert app.lbl_kpi_out_val.cget("text") == gt["avg_out_time"]
    assert app.lbl_kpi_hrs_val.cget("text") == gt["avg_working_hours"]

    # Verify that refreshing unfiltered dashboard does NOT recompute compute_time_series_metrics
    with patch("time_series_analysis.compute_time_series_metrics") as mock_calc:
        app._refresh_ts_dashboard()
        assert mock_calc.call_count == 0


# =========================================================================
# 3. Correct Business Unit, Department, Manager and Employee breakdowns
# =========================================================================
def test_breakdowns_matching_existing_analytical_engine(desktop_app):
    """Verify all 4 reporting levels match ground truth compute_level_breakdown results."""
    app = desktop_app
    app._reset_ts_filters()
    app.update()

    for level in ("Business Unit", "Department", "Reporting Manager", "Employee"):
        app._on_ts_level_changed(level)
        app.update()

        gt_breakdown = tsa.compute_level_breakdown(app.ts_dataset, level=level)
        rows = app._current_ts_breakdown_rows

        assert len(rows) == len(gt_breakdown)
        assert rows == gt_breakdown

        # Treeview rows count matches
        children = app.ts_tree.get_children()
        assert len(children) == len(gt_breakdown)


# =========================================================================
# 4. Search filtering prepared rows without recalculation
# =========================================================================
def test_search_filtering_without_recalculation(desktop_app):
    """Verify typing in search filters prepared rows in memory with 0 analytical recalculation."""
    app = desktop_app
    app._reset_ts_filters()
    app._on_ts_level_changed("Employee")
    app.update()
    assert len(app.ts_tree.get_children()) == 3

    # Patch analytical calculation methods
    with patch("time_series_analysis.compute_level_breakdown") as mock_bk, \
         patch("time_series_analysis.compute_time_series_metrics") as mock_m, \
         patch.object(snapshot_service, "get_breakdown") as mock_get_bk:

        # Search for "Alice"
        app.ts_search_var.set("Alice")
        app._render_ts_table_rows()  # direct render of debounced target
        app.update()

        children = app.ts_tree.get_children()
        assert len(children) == 1
        item_vals = app.ts_tree.item(children[0])["values"]
        assert "Alice" in str(item_vals[0])

        # Search for "E002" (Charlie's ID)
        app.ts_search_var.set("E002")
        app._render_ts_table_rows()
        app.update()

        children = app.ts_tree.get_children()
        assert len(children) == 1
        item_vals = app.ts_tree.item(children[0])["values"]
        assert "Charlie" in str(item_vals[0])

        # Clear search
        app.ts_search_var.set("")
        app._render_ts_table_rows()
        app.update()
        assert len(app.ts_tree.get_children()) == 3

        # Zero analytical recalculation calls occurred!
        assert mock_bk.call_count == 0
        assert mock_m.call_count == 0
        assert mock_get_bk.call_count == 0


# =========================================================================
# 5. Repeated reporting-level selection reusing prepared results
# =========================================================================
def test_repeated_reporting_level_selection_reusing_prepared_results(desktop_app):
    """Verify switching between reporting levels reuses prepared snapshot results."""
    app = desktop_app
    app._reset_ts_filters()
    app.update()

    with patch("time_series_analysis.compute_level_breakdown") as mock_calc:
        # Switch repeatedly between levels
        app._on_ts_level_changed("Business Unit")
        app.update()
        app._on_ts_level_changed("Department")
        app.update()
        app._on_ts_level_changed("Reporting Manager")
        app.update()
        app._on_ts_level_changed("Employee")
        app.update()
        app._on_ts_level_changed("Business Unit")
        app.update()

        # All 4 levels are precomputed at ingest; compute_level_breakdown must not be called
        assert mock_calc.call_count == 0


# =========================================================================
# 6. Correct dataset replacement
# =========================================================================
def test_correct_dataset_replacement(desktop_app, synthetic_excel, replacement_excel):
    """Verify loading a second dataset replaces the active snapshot and clears older results."""
    app = desktop_app

    # Load Dataset B (Replacement)
    snap2 = snapshot_service.prepare_dataset(str(replacement_excel))
    app._ts_load_success(snap2, str(replacement_excel), job_id=2)
    app.update()

    assert "Active: 1 records • 1 employees" in app.ts_lbl.cget("text")
    bu_vals = app.ts_bu_combo.cget("values")
    assert "Executive" in bu_vals
    assert "Engineering" not in bu_vals  # Old dataset options purged

    # Treeview and metrics update to new dataset
    assert len(app._current_ts_breakdown_rows) == 1
    assert app._current_ts_breakdown_rows[0]["Entity Name"] == "Executive"

    # Restore Dataset A for subsequent tests
    snap1 = snapshot_service.prepare_dataset(str(synthetic_excel))
    app._ts_load_success(snap1, str(synthetic_excel), job_id=3)
    app.update()
    assert "Active: 9 records • 3 employees" in app.ts_lbl.cget("text")


# =========================================================================
# 7. Failed upload preserving the existing dataset and UI state
# =========================================================================
def test_failed_upload_preserving_existing_dataset_and_ui_state(desktop_app, synthetic_excel, tmp_path):
    """Verify a failed upload preserves the previous active snapshot and UI status."""
    app = desktop_app

    initial_lbl_text = app.ts_lbl.cget("text")
    initial_rows = len(app._current_ts_breakdown_rows)
    initial_df = app.ts_dataset
    active_key = snapshot_service.get_active_snapshot().key

    # Attempt to upload a corrupt / invalid file
    bad_file = tmp_path / "corrupt_data.xlsx"
    bad_file.write_bytes(b"not an excel file")

    with patch("tkinter.messagebox.showerror") as mock_err:
        try:
            snapshot_service.prepare_dataset(str(bad_file))
        except Exception as e:
            app._ts_load_error(str(e), job_id=99)
        app.update()

        assert mock_err.called
        # Previous active snapshot must remain intact
        assert snapshot_service.has_active_snapshot()
        assert snapshot_service.get_active_snapshot().key == active_key
        assert app.ts_lbl.cget("text") == initial_lbl_text
        assert app.ts_dataset is initial_df
        assert len(app._current_ts_breakdown_rows) == initial_rows
        assert app.ts_file_path.get() == str(synthetic_excel)


# =========================================================================
# 8. Older processing jobs not overwriting newer results
# =========================================================================
def test_older_processing_jobs_not_overwriting_newer_results(desktop_app, synthetic_excel, replacement_excel):
    """Verify stale job fencing: Job 1 finishing after Job 2 does not overwrite Job 2."""
    app = desktop_app

    snap1 = snapshot_service.create_snapshot_for_testing(str(synthetic_excel)) if hasattr(snapshot_service, "create_snapshot_for_testing") else snapshot_service.prepare_dataset(str(synthetic_excel))
    snap2 = snapshot_service.prepare_dataset(str(replacement_excel))

    # Job 1 starts, Job 2 starts
    app._latest_ts_job_id = 50

    # Job 2 completes first with job_id=50
    app._ts_load_success(snap2, str(replacement_excel), job_id=50)
    app.update()
    assert "Active: 1 records • 1 employees" in app.ts_lbl.cget("text")

    # Job 1 completes later (stale callback with job_id=49)
    app._ts_load_success(snap1, str(synthetic_excel), job_id=49)
    app.update()

    # Job 1 completion must be rejected
    assert "Active: 1 records • 1 employees" in app.ts_lbl.cget("text")
    assert "Executive" in app.ts_bu_combo.cget("values")

    # Stale error callback must also be ignored
    with patch("tkinter.messagebox.showerror") as mock_err:
        app._ts_load_error("Some error from job 49", job_id=49)
        assert not mock_err.called

    # Restore Dataset A
    app._latest_ts_job_id = 51
    app._ts_load_success(snap1, str(synthetic_excel), job_id=51)
    app.update()


# =========================================================================
# 9. No Tkinter widget updates from background worker threads
# =========================================================================
def test_no_tkinter_widget_updates_from_background_worker_threads(desktop_app, synthetic_excel):
    """Verify background worker only executes analytical calculation and schedules UI updates."""
    app = desktop_app

    worker_thread_id = None
    after_called = False
    after_target = None

    def mock_after(ms, func):
        nonlocal after_called, after_target
        after_called = True
        after_target = func

    with patch.object(app, "after", side_effect=mock_after):
        def worker_wrapper():
            nonlocal worker_thread_id
            worker_thread_id = threading.get_ident()
            app._run_ts_load_job(str(synthetic_excel), job_id=100)

        t = threading.Thread(target=worker_wrapper)
        t.start()
        t.join()

        # Confirm worker thread was separate from main thread
        assert worker_thread_id != threading.get_ident()
        # Confirm app.after was invoked to schedule UI work on main thread
        assert after_called is True
        assert callable(after_target)


# =========================================================================
# 10. Excel export continuing to produce the same results and workbook structure
# =========================================================================
def test_excel_export_structure_and_values(desktop_app, tmp_path):
    """Verify Excel export produces valid workbook with expected sheets, metrics and breakdowns."""
    app = desktop_app
    app._reset_ts_filters()
    app.update()

    out_file = tmp_path / "desktop_export_test.xlsx"

    with patch("tkinter.filedialog.asksaveasfilename", return_value=str(out_file)), \
         patch("tkinter.messagebox.showinfo"):
        app._export_ts_excel()

    assert out_file.exists()

    # Verify sheet names match existing exporter structure
    wb = openpyxl.load_workbook(out_file, data_only=True)
    expected_sheets = ["KPI Summary", "Business Unit", "Department", "Reporting Manager", "Employee Breakdown"]
    assert wb.sheetnames == expected_sheets

    # Verify KPI Summary sheet contains expected metrics
    summary_ws = wb["KPI Summary"]
    summary_metrics_in_sheet = [
        row[1] for row in summary_ws.iter_rows(values_only=True) if row and len(row) > 1 and row[1] is not None
    ]
    assert "Total Records Analyzed" in summary_metrics_in_sheet
    assert "Unique Employees" in summary_metrics_in_sheet
    assert "Avg Days to Apply Leave (From Availed Date)" in summary_metrics_in_sheet
    assert "AVG Working Hours" in summary_metrics_in_sheet

    # Verify Breakdown sheets contain rows
    for s_name in ("Business Unit", "Department", "Reporting Manager", "Employee Breakdown"):
        ws = wb[s_name]
        data_rows = list(ws.iter_rows(values_only=True))
        assert len(data_rows) >= 2  # At least header + 1 data row
