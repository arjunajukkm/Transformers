import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

def generate_upload_file(input_file, output_file, employees):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objective Key Results"
    
    headers = [
        "Employee Number", "Employee Name", "Objective Title", "Objective Description", 
        "Time Frame", "Objective Type", "Objective Metric Type", "Objective Metric Name", 
        "Target Direction", "Objective Initial Value", "Objective Target Value", 
        "Objective Start Date", "Objective End Date", "Tags", "Include In Review", 
        "Visibility", "Progress Calculation Type", "Objective Weightage", 
        "KeyResult Employee Number", "KeyResult Employee Name", "Key Result Title", 
        "Key Result Description", "Key Result Metric Type", "Key Result Metric Name", 
        "Key Result Target Direction", "Key Result Initial Value", "Key Result Target Value", 
        "Key Result Start Date", "Key Result End Date", "Key Result Tags", "Key Result Weightage"
    ]
    
    ws.merge_cells('A1:AE1')
    title_cell = ws.cell(row=1, column=1, value="Objective Key Results")
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
    header_alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    ws.row_dimensions[2].height = 35
    
    xl = pd.ExcelFile(input_file)
    all_rows = []
    
    for sheet_name, employee_list in employees.items():
        actual_sheet_name = None
        for name in xl.sheet_names:
            if name.strip() == sheet_name.strip():
                actual_sheet_name = name
                break
                
        if not actual_sheet_name:
            continue
            
        df = pd.read_excel(xl, sheet_name=actual_sheet_name, header=0)
        extracted_data = []
        
        if df.shape[1] == 4:
            df.iloc[:, 0] = df.iloc[:, 0].ffill()
            df.iloc[:, 2] = df.iloc[:, 2].ffill()
            
            for kra, group in df.groupby(df.columns[0], sort=False):
                kra_weight = group.iloc[0, 2]
                kra_pct = kra_weight if pd.notna(kra_weight) else None
                
                weights = pd.to_numeric(group.iloc[:, 3], errors='coerce').fillna(0)
                total_w = weights.sum()
                
                for idx, row in group.iterrows():
                    kpi = row.iloc[1]
                    w = weights.loc[idx]
                    
                    if total_w > 0:
                        kpi_pct = round((w / total_w) * 100, 2)
                    else:
                        kpi_pct = 0.0
                        
                    kpi_pct_val = kpi_pct if pd.notna(kpi_pct) else None
                    extracted_data.append((kra, kpi, kpi_pct_val, kra_pct))
                    
        elif df.shape[1] >= 5:
            for _, row in df.iterrows():
                kra = row.iloc[0]
                kpi = row.iloc[1]
                kpi_pct = row.iloc[3]
                kra_pct = row.iloc[4]
                
                if pd.notna(kra) and pd.notna(kpi):
                    kpi_pct_val = kpi_pct if pd.notna(kpi_pct) else None
                    kra_pct_val = kra_pct if pd.notna(kra_pct) else None
                    extracted_data.append((kra, kpi, kpi_pct_val, kra_pct_val))
        
        for emp_num, emp_name in employee_list:
            for kra, kpi, kpi_pct, kra_pct in extracted_data:
                row_data = [
                    emp_num, emp_name, kra, None, "2025 - 26", "Individual", 
                    "Percentage", None, "Increase From", 0, 100, None, None, 
                    None, "Yes", "Managers", "Average of child KPI", kra_pct, 
                    emp_num, emp_name, kpi, None, "Percentage", None, 
                    "Increase From", 0, 100, None, None, None, kpi_pct
                ]
                all_rows.append((emp_num, row_data))
                
    data_alignment = Alignment(vertical="center", wrap_text=True)
    blue_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    row_font = Font(name="Calibri", size=10)
    
    current_emp = None
    use_blue = False
    
    current_row = 3
    for emp_num, row_data in all_rows:
        if emp_num != current_emp:
            current_emp = emp_num
            use_blue = not use_blue
            
        current_fill = blue_fill if use_blue else white_fill
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = data_alignment
            cell.fill = current_fill
            cell.font = row_font
            
        ws.row_dimensions[current_row].height = 30
        current_row += 1
        
    ws.freeze_panes = "A3"
    
    col_widths = {
        1: 15, 2: 25, 3: 70, 4: 12, 5: 12, 6: 12, 7: 18, 8: 12, 9: 12, 10: 12, 
        11: 12, 12: 12, 13: 12, 14: 12, 15: 12, 16: 12, 17: 12, 18: 18, 19: 15, 
        20: 25, 21: 70, 22: 12, 23: 18, 24: 12, 25: 12, 26: 12, 27: 12, 28: 12, 
        29: 12, 30: 12, 31: 18
    }
    
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        
    wb.save(output_file)


def normalize_okr_weightage(val):
    """Normalize OKR source weightage.
    
    Rule:
      if numeric_weight <= 1:
          objective_weightage = numeric_weight * 100
      else:
          objective_weightage = numeric_weight
    Whole numbers have redundant decimal formatting removed (e.g. 50.0 -> 50).
    """
    if val is None or pd.isna(val):
        return None
    s = str(val).strip()
    if not s:
        return None
    if s.endswith('%'):
        s = s[:-1].strip()
    try:
        num = float(s)
    except (ValueError, TypeError):
        return val
        
    if num <= 1.0:
        num = num * 100.0
        
    if round(num, 6).is_integer():
        return int(round(num))
    else:
        return round(num, 2)


def _detect_okr_columns(columns):
    """Detect matching column names for OKR, Key Result Areas, and Weightage.
    
    Recognizes standard synonyms:
    - OKR: 'OKR', 'OKRs', 'Objective', 'Objectives', 'Objective/Goal', 'Goal', 'Parent Objective'
    - KR: 'Key Result Areas', 'Key Results (KRs)', 'Key Results', 'Key Result', 'KR', 'KRs', 'KPI'
    - Weightage: 'Weightage', 'Weight', 'Weight (%)', 'Weightage (%)', 'Weights', 'Objective Weightage'
    """
    col_map = {c: str(c).strip().lower() for c in columns if pd.notna(c) and str(c).strip()}
    
    okr_col = None
    kr_col = None
    weight_col = None
    
    # 1. OKR / Objective / Goal
    for orig, s in col_map.items():
        if any(ex in s for ex in ['key result', 'weight', 'metric', 'type', 'cadence', 'frequency', 'target direction']):
            continue
        if (s == 'okr' or s.startswith('okr') or 'okr' in s.split() or 
            'objective' in s or 'goal' in s or s in ['parent objective', 'kra']):
            okr_col = orig
            break
            
    # 2. Key Result / KR / KPI
    for orig, s in col_map.items():
        if orig == okr_col or 'weight' in s:
            continue
        if ('key result' in s or s in ['kr', 'krs', 'kpi', 'kpis', 'key results'] or 
            '(kr' in s or 'key result areas' in s or 'key result area' in s):
            kr_col = orig
            break
            
    # 3. Weightage / Weight
    for orig, s in col_map.items():
        if orig in (okr_col, kr_col):
            continue
        if 'weight' in s or s == 'wt' or s.startswith('weight'):
            weight_col = orig
            break
            
    return okr_col, kr_col, weight_col


def _find_okr_columns(df):
    """Find and return the actual column names for OKR, Key Result Areas, and Weightage.
    
    Headers are normalized by trimming whitespace, case-insensitivity, and recognizing
    standard synonyms.
    Raises ValueError if any required column is missing.
    """
    okr_col, kr_col, weight_col = _detect_okr_columns(df.columns)
                
    missing = []
    if okr_col is None:
        missing.append("OKR")
    if kr_col is None:
        missing.append("Key Result Areas")
    if weight_col is None:
        missing.append("Weightage")
        
    if missing:
        if len(missing) == 1:
            raise ValueError(f"The uploaded OKR file is missing the required column: {missing[0]}.")
        else:
            raise ValueError(f"The uploaded OKR file is missing required columns: {', '.join(missing)}.")
            
    return okr_col, kr_col, weight_col


def _parse_okr_dataframe(df):
    """Parse hierarchical OKR structure into a list of (objective_title, kr_title, objective_weightage).
    
    Rules:
    - OKR is provided on the first row of an Objective block.
    - Subsequent rows until next OKR belong to the current Objective.
    - Objective Weightage is determined by the first non-empty Weightage value belonging to that Objective.
    - Weightage is applied to every output row generated for that Objective.
    - Do NOT carry Weightage from one Objective into a different Objective when the new Objective itself has no Weightage.
    - Ignore completely blank rows.
    - Do NOT generate a row where Key Result Areas is blank.
    - Automatically detects headers if located within the first 15 rows.
    """
    # If headers are not in df.columns, check if a header row exists in the first 15 rows
    okr_col, kr_col, weight_col = _detect_okr_columns(df.columns)
    if not (okr_col and kr_col and weight_col):
        for r in range(min(15, len(df))):
            row_vals = df.iloc[r].dropna().tolist()
            o, k, w = _detect_okr_columns(row_vals)
            if o and k and w:
                new_df = df.iloc[r + 1:].copy()
                new_df.columns = df.iloc[r].tolist()
                df = new_df
                break

    okr_col, kr_col, weight_col = _find_okr_columns(df)
    
    objective_blocks = []
    current_block = None
    
    for _, row in df.iterrows():
        raw_okr = row[okr_col]
        raw_kr = row[kr_col]
        raw_wt = row[weight_col]
        
        has_okr = pd.notna(raw_okr) and str(raw_okr).strip() != ""
        has_kr = pd.notna(raw_kr) and str(raw_kr).strip() != ""
        has_wt = pd.notna(raw_wt) and str(raw_wt).strip() != ""
        
        if not (has_okr or has_kr or has_wt):
            continue
            
        if has_okr:
            current_block = {
                "objective": str(raw_okr).strip(),
                "rows": []
            }
            objective_blocks.append(current_block)
            
        if current_block is not None:
            current_block["rows"].append({
                "kr": str(raw_kr).strip() if has_kr else None,
                "weight": raw_wt if has_wt else None
            })
            
    if not objective_blocks:
        raise ValueError("No valid Objectives were found in the uploaded file.")
        
    extracted_data = []
    for block in objective_blocks:
        obj_title = block["objective"]
        first_weight = None
        for r in block["rows"]:
            if r["weight"] is not None:
                first_weight = r["weight"]
                break
                
        obj_weight = normalize_okr_weightage(first_weight) if first_weight is not None else None
        
        for r in block["rows"]:
            kr_title = r["kr"]
            if kr_title:
                extracted_data.append((obj_title, kr_title, obj_weight))
                
    if not extracted_data:
        raise ValueError("No valid Key Result Areas were found in the uploaded file.")
        
    return extracted_data


def generate_okr_upload_file(input_file, output_file, employees):
    """Generate KEKA Objective Key Results Excel file for OKR Upload mode.
    
    Output contains exactly 30 headers (A to AD), with merged title A1:AD1,
    and no Key Result Weightage column.
    """
    if not os.path.exists(input_file):
        raise FileNotFoundError(f"Input file not found: {input_file}")
        
    if not employees:
        raise ValueError("At least one employee mapping must be provided.")
        
    xl = pd.ExcelFile(input_file)
    available_sheets = xl.sheet_names
    if not available_sheets:
        raise ValueError("The uploaded Excel workbook contains no sheets.")
        
    is_single_sheet = len(available_sheets) == 1
    single_sheet_name = available_sheets[0] if is_single_sheet else None
    
    sheet_extracted_cache = {}
    all_rows = []
    matched_any_sheet = False
    
    for req_sheet_name, employee_list in employees.items():
        if not employee_list:
            continue
            
        target_sheet = None
        if is_single_sheet:
            target_sheet = single_sheet_name
        else:
            for s in available_sheets:
                if s.strip().lower() == req_sheet_name.strip().lower():
                    target_sheet = s
                    break
                    
        if not target_sheet:
            continue
            
        matched_any_sheet = True
        
        if target_sheet not in sheet_extracted_cache:
            df = pd.read_excel(xl, sheet_name=target_sheet, header=0)
            sheet_extracted_cache[target_sheet] = _parse_okr_dataframe(df)
            
        extracted_data = sheet_extracted_cache[target_sheet]
        
        for emp_num, emp_name in employee_list:
            for obj_title, kr_title, obj_weight in extracted_data:
                row_data = [
                    emp_num, emp_name, obj_title, None, "2025 - 26", "Individual", 
                    "Percentage", None, "Increase From", 0, 100, None, None, 
                    None, "Yes", "Managers", "Average of child KPI", obj_weight, 
                    emp_num, emp_name, kr_title, None, "Percentage", None, 
                    "Increase From", 0, 100, None, None, None
                ]
                all_rows.append((emp_num, row_data))
                
    if not matched_any_sheet:
        available_str = ", ".join(f"'{s}'" for s in available_sheets)
        raise ValueError(
            f"None of the mapped designations/sheets matched the sheets in the uploaded file.\n"
            f"Available sheets: {available_str}"
        )
        
    if not all_rows:
        raise ValueError("No output rows could be generated from the mapped employees and OKR data.")
        
    wb = Workbook()
    ws = wb.active
    ws.title = "Objective Key Results"
    
    headers = [
        "Employee Number", "Employee Name", "Objective Title", "Objective Description", 
        "Time Frame", "Objective Type", "Objective Metric Type", "Objective Metric Name", 
        "Target Direction", "Objective Initial Value", "Objective Target Value", 
        "Objective Start Date", "Objective End Date", "Tags", "Include In Review", 
        "Visibility", "Progress Calculation Type", "Objective Weightage", 
        "KeyResult Employee Number", "KeyResult Employee Name", "Key Result Title", 
        "Key Result Description", "Key Result Metric Type", "Key Result Metric Name", 
        "Key Result Target Direction", "Key Result Initial Value", "Key Result Target Value", 
        "Key Result Start Date", "Key Result End Date", "Key Result Tags"
    ]
    
    # Exactly 30 columns: A1:AD1
    ws.merge_cells('A1:AD1')
    title_cell = ws.cell(row=1, column=1, value="Objective Key Results")
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 25
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
    header_alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    ws.row_dimensions[2].height = 35
    
    data_alignment = Alignment(vertical="center", wrap_text=True)
    blue_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    row_font = Font(name="Calibri", size=10)
    
    current_emp = None
    use_blue = False
    
    current_row = 3
    for emp_num, row_data in all_rows:
        if emp_num != current_emp:
            current_emp = emp_num
            use_blue = not use_blue
            
        current_fill = blue_fill if use_blue else white_fill
        
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.alignment = data_alignment
            cell.fill = current_fill
            cell.font = row_font
            
        ws.row_dimensions[current_row].height = 30
        current_row += 1
        
    ws.freeze_panes = "A3"
    
    col_widths = {
        1: 15, 2: 25, 3: 70, 4: 12, 5: 12, 6: 12, 7: 18, 8: 12, 9: 12, 10: 12, 
        11: 12, 12: 12, 13: 12, 14: 12, 15: 12, 16: 12, 17: 12, 18: 18, 19: 15, 
        20: 25, 21: 70, 22: 12, 23: 18, 24: 12, 25: 12, 26: 12, 27: 12, 28: 12, 
        29: 12, 30: 12
    }
    
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        
    wb.save(output_file)


if __name__ == "__main__":
    input_file = "YOUR_INPUT_FILE.xlsx"
    output_file = "YOUR_OUTPUT_FILE.xlsx"

    employees = {
        "Sheet Name Exactly As In Excel": [
            ("EMP001", "Employee Full Name"),
            ("EMP002", "Another Employee"),
        ],
        "Another Sheet Name": [
            ("EMP003", "Third Employee"),
        ],
    }

    generate_upload_file(input_file, output_file, employees)
