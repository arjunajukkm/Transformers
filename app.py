import threading
import traceback
import sys
import os
import ctypes
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import pandas as pd
import customtkinter as ctk
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

import ui_components as ui
import time_series_analysis as tsa
from storage import snapshot_service

ctk.set_appearance_mode("Dark")
ctk.set_default_color_theme("blue")

APP_TITLE = "Transformers"
WINDOW_SIZE = "1280x800"

class App(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title(APP_TITLE)
        self.geometry(WINDOW_SIZE)
        self.minsize(1024, 700)

        try:
            myappid = 'moshpit.transformers.1.0'
            ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(myappid)
        except Exception:
            pass

        try:
            if getattr(sys, 'frozen', False):
                base_path = Path(sys._MEIPASS)
            else:
                base_path = Path(__file__).parent
            icon_path = base_path / "icon.ico"
            if icon_path.exists():
                self.iconbitmap(str(icon_path))
        except Exception as e:
            pass

        # StringVars for file paths
        self.selected_input_file = ctk.StringVar()
        self.selected_gen_input_file = ctk.StringVar()
        self.selected_okr_input_file = ctk.StringVar()
        self.gen_mode_var = ctk.StringVar(value="Existing KRA Upload")
        self.absent_emp_file = ctk.StringVar()
        self.absent_att_file = ctk.StringVar()
        self.absent_wfh_file = ctk.StringVar()
        self.att_summary_file = ctk.StringVar()
        self.analyse_upload_file = ctk.StringVar()

        # Time Series Analysis state
        self.ts_dataset = None
        self.ts_file_path = ctk.StringVar()
        self.ts_bu_var = ctk.StringVar(value="All Business Units")
        self.ts_dept_var = ctk.StringVar(value="All Departments")
        self.ts_month_var = ctk.StringVar(value="All Months")
        self.ts_level_var = ctk.StringVar(value="Business Unit")
        self.ts_search_var = ctk.StringVar()
        self.is_ts_loading = False
        self._latest_ts_job_id = 0
        self._current_ts_breakdown_rows = []
        self._search_debounce_id = None

        # Time and Leave Master multi-file lists
        self.tl_perf_files = []
        self.tl_leave_active_files = []
        self.tl_leave_inactive_files = []
        self.tl_wfh_files = []
        
        # State flags
        self.is_processing = False
        self.is_generating = False
        self.is_analysing = False
        
        # Navigation state
        self.transform_menu_expanded = False
        self.analyse_menu_expanded = False
        self.menu_expanded = False
        self.sidebar_collapsed = False
        self.current_active_frame = "dashboard"
        self.nav_parent_btn = None
        self.nav_chevron = None
        self.analyse_parent_btn = None
        self.analyse_chevron = None
        self.sub_nav_btns = {}

        self._setup_window()
        self._build_layout()
        self.select_frame_by_name("dashboard")

    def _setup_window(self):
        self.grid_rowconfigure(0, weight=1)
        self.grid_columnconfigure(0, weight=0)
        self.grid_columnconfigure(1, weight=1)
        self.configure(fg_color=ui.COLOR_BG)

    def _build_layout(self):
        # Sidebar
        self.sidebar = ctk.CTkFrame(self, width=280, corner_radius=0, fg_color=ui.COLOR_SIDEBAR)
        self.sidebar.grid(row=0, column=0, sticky="nsew")
        self.sidebar.grid_propagate(False)
        self.sidebar.grid_rowconfigure(10, weight=1)
        self.sidebar.grid_columnconfigure(0, weight=1)

        # Header with Logo & Collapse Toggle Bar (Row 0)
        self.sidebar_header = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        self.sidebar_header.grid(row=0, column=0, sticky="ew", padx=14, pady=(24, 20))
        self.sidebar_header.grid_columnconfigure(0, weight=1)
        self.sidebar_header.grid_columnconfigure(1, weight=0)

        self.lbl_logo = ctk.CTkLabel(
            self.sidebar_header, text="Transformers",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=24, weight="bold"),
            text_color=ui.COLOR_ACCENT, anchor="w",
        )
        self.lbl_logo.grid(row=0, column=0, sticky="w", padx=(6, 0))

        self.lbl_logo_compact = ctk.CTkLabel(
            self.sidebar_header, text="⚡",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=22, weight="bold"),
            text_color=ui.COLOR_ACCENT, anchor="center",
        )

        self.btn_sidebar_toggle = ctk.CTkButton(
            self.sidebar_header, text="◀", width=28, height=28, corner_radius=6,
            fg_color=ui.COLOR_CARD, hover_color=ui.COLOR_NAV_HOVER,
            text_color=ui.COLOR_TEXT_SEC, font=ctk.CTkFont(size=12, weight="bold"),
            command=self.toggle_sidebar,
        )
        self.btn_sidebar_toggle.grid(row=0, column=1, sticky="e")
        ui.create_tooltip(self.btn_sidebar_toggle, "Collapse Sidebar (expand workspace)")

        sep = ctk.CTkFrame(self.sidebar, height=1, fg_color=ui.COLOR_SIDEBAR_SEP)
        sep.grid(row=1, column=0, sticky="ew", padx=14, pady=(0, 16))

        # ── Collapsible Transform Navigation Group (Closed by Default) ────────
        self.nav_parent_frame, self.nav_parent_btn, self.nav_chevron = ui.create_collapsible_nav_item(
            self.sidebar,
            text="Transform",
            icon=ui.ICON_HOME,
            on_select=lambda: self._on_transform_parent_clicked(),
            on_toggle=lambda: self.toggle_transform_menu(),
            row=2,
        )
        self.nav_chevron.configure(text="▸")

        # Sub-menu container for child modules (Starts Closed)
        self.sub_menu_frame = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        self.sub_menu_frame.grid_columnconfigure(0, weight=1)

        # Sub-items: KRA Management, Absent Management, Attendance Summary, Time and Leave Master
        transform_sub_items = [
            ("transform", "KRA Management", ui.ICON_KRA),
            ("absent", "Absent Management", ui.ICON_ABSENT),
            ("att_summary", "Attendance Summary", ui.ICON_ATTENDANCE),
            ("time_leave", "Time and Leave Master", ui.ICON_TIME_LEAVE),
        ]

        for idx, (name, text, icon) in enumerate(transform_sub_items):
            btn = ui.create_sub_nav_button(
                self.sub_menu_frame,
                text=text,
                icon=icon,
                command=lambda n=name: self.select_frame_by_name(n),
                row=idx,
            )
            self.sub_nav_btns[name] = btn

        # ── Collapsible Analyse Navigation Group (Closed by Default) ──────────
        self.analyse_parent_frame, self.analyse_parent_btn, self.analyse_chevron = ui.create_collapsible_nav_item(
            self.sidebar,
            text="Analyse",
            icon=ui.ICON_ANALYSE,
            on_select=lambda: self._on_analyse_parent_clicked(),
            on_toggle=lambda: self.toggle_analyse_menu(),
            row=4,
        )
        self.analyse_chevron.configure(text="▸")

        # Sub-menu container for Analyse child modules (Starts Closed)
        self.analyse_sub_menu_frame = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        self.analyse_sub_menu_frame.grid_columnconfigure(0, weight=1)

        # Sub-items: Workforce Intelligence, Time Series Analysis, Upload
        analyse_sub_items = [
            ("workforce_intelligence", "Workforce Intelligence", ui.ICON_ANALYSE),
            ("analyse_time_series", "Time Series Analysis", ui.ICON_DASHBOARD),
            ("analyse_upload", "Upload", ui.ICON_UPLOAD),
        ]

        for idx, (name, text, icon) in enumerate(analyse_sub_items):
            btn = ui.create_sub_nav_button(
                self.analyse_sub_menu_frame,
                text=text,
                icon=icon,
                command=lambda n=name: self.select_frame_by_name(n),
                row=idx,
            )
            self.sub_nav_btns[name] = btn

        # Main Content (Takes all remaining width)
        self.main_content = ctk.CTkFrame(self, fg_color="transparent", corner_radius=0)
        self.main_content.grid(row=0, column=1, sticky="nsew", padx=30, pady=30)
        self.main_content.grid_rowconfigure(0, weight=1)
        self.main_content.grid_columnconfigure(0, weight=1)

        # Build Frames
        self.frames = {}
        self._build_dashboard_frame()
        self._build_transform_frame()
        self._build_generate_frame()
        self._build_absent_frame()
        self._build_att_summary_frame()
        self._build_time_leave_frame()
        self._build_time_series_frame()
        self._build_analyse_upload_frame()
        self._build_workforce_intelligence_frame()

    def toggle_sidebar(self, force_state=None):
        """Toggle sidebar between expanded (280px) and collapsed icon-rail (60px)."""
        if force_state is not None:
            self.sidebar_collapsed = force_state
        else:
            self.sidebar_collapsed = not self.sidebar_collapsed

        if self.sidebar_collapsed:
            # 1. Collapse sidebar container to 60px
            self.sidebar.configure(width=60)
            self.sidebar.grid_propagate(False)

            # 2. Compact header
            self.lbl_logo.grid_remove()
            self.lbl_logo_compact.grid(row=0, column=0, sticky="nsew", padx=0)
            self.btn_sidebar_toggle.grid(row=0, column=1, sticky="e", padx=(0, 4))
            self.btn_sidebar_toggle.configure(text="▶", width=22, height=22)
            ui.create_tooltip(self.btn_sidebar_toggle, "Expand Sidebar")

            # 3. Hide chevrons and child submenus
            self.nav_chevron.grid_remove()
            self.analyse_chevron.grid_remove()
            self.sub_menu_frame.grid_remove()
            self.analyse_sub_menu_frame.grid_remove()

            # 4. Icon-only parent buttons
            self.nav_parent_frame.grid_configure(padx=6)
            self.analyse_parent_frame.grid_configure(padx=6)
            self.nav_parent_btn.configure(text=ui.ICON_HOME, anchor="center")
            self.analyse_parent_btn.configure(text=ui.ICON_ANALYSE, anchor="center")
            ui.create_tooltip(self.nav_parent_btn, "Transform (click to expand)")
            ui.create_tooltip(self.analyse_parent_btn, "Analyse (click to expand)")
        else:
            # 1. Expand sidebar container to 280px
            self.sidebar.configure(width=280)
            self.sidebar.grid_propagate(False)

            # 2. Expanded header
            self.lbl_logo_compact.grid_remove()
            self.lbl_logo.grid(row=0, column=0, sticky="w", padx=(6, 0))
            self.btn_sidebar_toggle.grid(row=0, column=1, sticky="e", padx=0)
            self.btn_sidebar_toggle.configure(text="◀", width=28, height=28)
            ui.create_tooltip(self.btn_sidebar_toggle, "Collapse Sidebar (expand workspace)")

            # 3. Restore parent buttons
            self.nav_parent_frame.grid_configure(padx=14)
            self.analyse_parent_frame.grid_configure(padx=14)
            self.nav_parent_btn.configure(text=f"  {ui.ICON_HOME}   Transform", anchor="w")
            self.analyse_parent_btn.configure(text=f"  {ui.ICON_ANALYSE}   Analyse", anchor="w")
            self.nav_chevron.grid()
            self.analyse_chevron.grid()

            # 4. Restore submenus to their previous expansion states
            if self.transform_menu_expanded:
                self.sub_menu_frame.grid(row=3, column=0, sticky="ew")
                self.nav_chevron.configure(text="▾")
            else:
                self.sub_menu_frame.grid_remove()
                self.nav_chevron.configure(text="▸")

            if self.analyse_menu_expanded:
                self.analyse_sub_menu_frame.grid(row=5, column=0, sticky="ew")
                self.analyse_chevron.configure(text="▾")
            else:
                self.analyse_sub_menu_frame.grid_remove()
                self.analyse_chevron.configure(text="▸")

        self.update_idletasks()

    def _on_transform_parent_clicked(self):
        """When Transform parent is clicked, toggle submenu and show overview."""
        if self.sidebar_collapsed:
            self.toggle_sidebar(force_state=False)
            self.toggle_transform_menu(force_state=True)
        else:
            self.toggle_transform_menu()
        self.select_frame_by_name("dashboard")

    def toggle_transform_menu(self, force_state=None):
        """Toggle the collapsible Transform sub-menu between expanded and collapsed states."""
        if force_state is not None:
            self.transform_menu_expanded = force_state
        else:
            self.transform_menu_expanded = not self.transform_menu_expanded
        self.menu_expanded = self.transform_menu_expanded

        if not self.sidebar_collapsed:
            if self.transform_menu_expanded:
                self.sub_menu_frame.grid(row=3, column=0, sticky="ew")
                self.nav_chevron.configure(text="▾")
            else:
                self.sub_menu_frame.grid_remove()
                self.nav_chevron.configure(text="▸")
        # Update active indicator on parent if needed
        self._refresh_nav_active_states()

    def _on_analyse_parent_clicked(self):
        """When Analyse parent is clicked, toggle submenu and show active/default frame."""
        if self.sidebar_collapsed:
            self.toggle_sidebar(force_state=False)
            self.toggle_analyse_menu(force_state=True)
        else:
            self.toggle_analyse_menu()

        cur = getattr(self, "current_active_frame", None)
        if cur not in ("workforce_intelligence", "analyse_time_series", "analyse_upload"):
            self.select_frame_by_name("workforce_intelligence")
        else:
            self.select_frame_by_name(cur)

    def toggle_analyse_menu(self, force_state=None):
        """Toggle the collapsible Analyse sub-menu between expanded and collapsed states."""
        if force_state is not None:
            self.analyse_menu_expanded = force_state
        else:
            self.analyse_menu_expanded = not self.analyse_menu_expanded

        if not self.sidebar_collapsed:
            if self.analyse_menu_expanded:
                self.analyse_sub_menu_frame.grid(row=5, column=0, sticky="ew")
                self.analyse_chevron.configure(text="▾")
            else:
                self.analyse_sub_menu_frame.grid_remove()
                self.analyse_chevron.configure(text="▸")
        # Update active indicator on parent if needed
        self._refresh_nav_active_states()

    def _refresh_nav_active_states(self):
        """Update parent and child navigation highlight based on current screen and submenu states."""
        name = getattr(self, "current_active_frame", "dashboard")
        transform_children = ("dashboard", "transform", "generate", "absent", "att_summary", "time_leave")
        analyse_children = ("workforce_intelligence", "analyse_time_series", "analyse_dashboard", "analyse_upload")

        is_transform_parent = (name == "dashboard") or (name in transform_children and not self.transform_menu_expanded)
        is_analyse_parent = (name in analyse_children and not self.analyse_menu_expanded)

        ui.set_nav_active(self.nav_parent_btn, is_transform_parent)
        ui.set_nav_active(self.analyse_parent_btn, is_analyse_parent)

        for n, btn in self.sub_nav_btns.items():
            is_active = (n == name) or (name == "generate" and n == "transform")
            ui.set_sub_nav_active(btn, is_active)

    def select_frame_by_name(self, name: str):
        if name == "analyse_dashboard":
            name = "analyse_time_series"
        self.current_active_frame = name

        # Do NOT auto-expand menus on screen selection!
        # Just update parent and child highlight based on whether the submenu is expanded or collapsed
        self._refresh_nav_active_states()

        # Hide all frames
        for f in self.frames.values():
            f.grid_forget()

        # Show selected
        if name in self.frames:
            self.frames[name].grid(row=0, column=0, sticky="nsew")
            if name == "workforce_intelligence" and hasattr(self, "workforce_dashboard_view"):
                self.workforce_dashboard_view.sync_snapshot_state()


    # ---------------------------------------------------------
    # Transform Hub / Overview
    # ---------------------------------------------------------
    def _build_dashboard_frame(self):
        f = ctk.CTkFrame(self.main_content, fg_color="transparent")
        self.frames["dashboard"] = f
        f.grid_columnconfigure((0, 1), weight=1)

        ui.create_page_header(f, "Transform", "Quick access to all data transformation and processing modules.").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 30))

        ui.create_dashboard_card(f, ui.ICON_KRA, "KRA Management", "Transform KRA Excel sheets and generate final KEKA upload files.", "Active", lambda: self.select_frame_by_name("transform"), 1, 0)
        ui.create_dashboard_card(f, ui.ICON_ABSENT, "Absent Management", "Process employee data to generate Absent Intimation Reports.", "Active", lambda: self.select_frame_by_name("absent"), 1, 1)
        ui.create_dashboard_card(f, ui.ICON_ATTENDANCE, "Attendance Summary", "Generate summarized attendance reports from raw portal data.", "Active", lambda: self.select_frame_by_name("att_summary"), 2, 0)
        ui.create_dashboard_card(f, ui.ICON_TIME_LEAVE, "Time and Leave Master", "Manage and process time and leave records.", "Active", lambda: self.select_frame_by_name("time_leave"), 2, 1)

    # ---------------------------------------------------------
    # KRA Management
    # ---------------------------------------------------------
    def _build_transform_frame(self):
        f = ctk.CTkFrame(self.main_content, fg_color="transparent")
        self.frames["transform"] = f
        f.grid_columnconfigure(0, weight=1)

        hdr = ui.create_page_header(f, "KRA Management", "Transform raw KRA files and generate system uploads.")
        hdr.grid(row=0, column=0, sticky="ew", pady=(0, 20))
        
        # Actions wrapper
        actions = ctk.CTkFrame(hdr, fg_color="transparent")
        actions.grid(row=0, column=1, sticky="e")
        ui.create_secondary_button(actions, "Go to Generation →", lambda: self.select_frame_by_name("generate"), 150).pack()

        card = ui.create_card(f)
        card.grid(row=1, column=0, sticky="nsew")

        ctk.CTkLabel(card, text="1. Transform Raw File", font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"), text_color=ui.COLOR_TEXT).grid(row=0, column=0, sticky="w", padx=16, pady=(16, 4))
        ctk.CTkLabel(card, text="Upload raw Excel sheet to structure KRA/KPI data.", font=ctk.CTkFont(family=ui.FONT_FAMILY, size=12), text_color=ui.COLOR_TEXT_SEC).grid(row=1, column=0, sticky="w", padx=16, pady=(0, 12))

        ui.create_upload_row(card, self.selected_input_file, "Select Raw KRA Excel...", "Browse", 2)

        self.kra_prog = ctk.CTkProgressBar(card, mode="indeterminate", height=4, corner_radius=2, fg_color=ui.COLOR_INPUT_BG, progress_color=ui.COLOR_ACCENT)
        self.kra_prog.grid(row=3, column=0, sticky="ew", padx=16, pady=(8, 0))
        self.kra_prog.set(0)
        self.kra_prog.grid_remove()

        row4 = ctk.CTkFrame(card, fg_color="transparent")
        row4.grid(row=4, column=0, sticky="ew", padx=16, pady=(12, 16))
        
        _, self.kra_dot, self.kra_lbl = ui.create_status_badge(row4, "Ready")
        self.kra_lbl.master.pack(side="left")

        self.btn_kra_tf = ui.create_primary_button(row4, "Transform File", self.start_transform)
        self.btn_kra_tf.pack(side="right")

    def _build_generate_frame(self):
        f = ctk.CTkFrame(self.main_content, fg_color="transparent")
        self.frames["generate"] = f
        f.grid_columnconfigure(0, weight=1)

        hdr = ui.create_page_header(f, "Generate Upload", "Create final KEKA upload file using employee mappings.")
        hdr.grid(row=0, column=0, sticky="ew", pady=(0, 20))
        
        actions = ctk.CTkFrame(hdr, fg_color="transparent")
        actions.grid(row=0, column=1, sticky="e")
        ui.create_secondary_button(actions, "← Back to Transform", lambda: self.select_frame_by_name("transform"), 150).pack()

        card = ui.create_card(f)
        card.grid(row=1, column=0, sticky="nsew")

        # Mode Selector
        mode_row = ctk.CTkFrame(card, fg_color="transparent")
        mode_row.grid(row=0, column=0, sticky="w", padx=16, pady=(16, 10))

        ctk.CTkLabel(
            mode_row, text="Generation Mode:",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=13, weight="bold"),
            text_color=ui.COLOR_TEXT_SEC
        ).pack(side="left", padx=(0, 12))

        self.gen_mode_seg = ctk.CTkSegmentedButton(
            mode_row,
            values=["Existing KRA Upload", "OKR Upload"],
            command=self._on_gen_mode_changed,
            selected_color=ui.COLOR_ACCENT,
            selected_hover_color=ui.COLOR_ACCENT_HOVER,
            unselected_color=ui.COLOR_INPUT_BG,
            unselected_hover_color=ui.COLOR_BTN_SEC_HOV,
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=12, weight="bold"),
        )
        self.gen_mode_seg.set("Existing KRA Upload")
        self.gen_mode_seg.pack(side="left")

        # Dynamic Section Header Label
        self.gen_upload_title_lbl = ctk.CTkLabel(
            card, text="Upload Transformed File",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"),
            text_color=ui.COLOR_TEXT
        )
        self.gen_upload_title_lbl.grid(row=1, column=0, sticky="w", padx=16, pady=(4, 4))

        # Upload controls for each mode
        self.gen_kra_upload_frame, _, _ = ui.create_upload_row(
            card, self.selected_gen_input_file, "Select transformed Excel...", "Browse", 2
        )
        self.gen_okr_upload_frame, _, _ = ui.create_upload_row(
            card, self.selected_okr_input_file, "Expected columns: OKR, Key Result Areas, Weightage", "Browse", 2
        )
        self.gen_okr_upload_frame.grid_remove()

        ctk.CTkLabel(
            card, text="Employee Mapping Format:\nEMP ID, Name, Designation (Sheet Name)",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=13, weight="bold"),
            text_color=ui.COLOR_TEXT, justify="left"
        ).grid(row=3, column=0, sticky="w", padx=16, pady=(12, 8))
        self.gen_mapping_text = ctk.CTkTextbox(
            card, height=80, font=ctk.CTkFont(family="Consolas", size=12),
            fg_color=ui.COLOR_INPUT_BG, border_color=ui.COLOR_BORDER, border_width=1,
            text_color=ui.COLOR_TEXT
        )
        self.gen_mapping_text.grid(row=4, column=0, sticky="ew", padx=16, pady=(0, 12))
        self.gen_mapping_text.insert("0.0", 'EMP001, John Doe, Manager\nEMP002, Jane Smith, Manager')

        self.gen_prog = ctk.CTkProgressBar(card, mode="indeterminate", height=4, corner_radius=2, fg_color=ui.COLOR_INPUT_BG, progress_color=ui.COLOR_ACCENT)
        self.gen_prog.grid(row=5, column=0, sticky="ew", padx=16, pady=(4, 0))
        self.gen_prog.set(0)
        self.gen_prog.grid_remove()

        row6 = ctk.CTkFrame(card, fg_color="transparent")
        row6.grid(row=6, column=0, sticky="ew", padx=16, pady=(12, 16))
        
        _, self.gen_dot, self.gen_lbl = ui.create_status_badge(row6, "Ready")
        self.gen_lbl.master.pack(side="left")

        self.btn_gen = ui.create_primary_button(row6, "Generate File", self.start_generate)
        self.btn_gen.pack(side="right")

    def _on_gen_mode_changed(self, mode: str):
        if mode == "OKR Upload":
            self.gen_upload_title_lbl.configure(text="Upload OKR Source File")
            self.gen_kra_upload_frame.grid_remove()
            self.gen_okr_upload_frame.grid()
        else:
            self.gen_upload_title_lbl.configure(text="Upload Transformed File")
            self.gen_okr_upload_frame.grid_remove()
            self.gen_kra_upload_frame.grid()

    # ---------------------------------------------------------
    # Absent Management
    # ---------------------------------------------------------
    def _build_absent_frame(self):
        f = ctk.CTkFrame(self.main_content, fg_color="transparent")
        self.frames["absent"] = f
        f.grid_columnconfigure(0, weight=1)

        ui.create_page_header(f, "Absent Management", "Generate absent intimation reports by comparing master, attendance and OD/WFH data.").grid(row=0, column=0, sticky="ew", pady=(0, 20))

        card = ui.create_card(f)
        card.grid(row=1, column=0, sticky="nsew")
        
        ctk.CTkLabel(card, text="Source Files", font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"), text_color=ui.COLOR_TEXT).grid(row=0, column=0, sticky="w", padx=16, pady=(16, 12))

        ui.create_upload_row(card, self.absent_emp_file, "1. Employee Master (Expected: Employee Master.xlsx)", "Browse", 1)
        ui.create_upload_row(card, self.absent_att_file, "2. Attendance Report (Expected: Attendance Report.xlsx)", "Browse", 2)
        ui.create_upload_row(card, self.absent_wfh_file, "3. OD/WFH Application Report (Expected: OD_WFH Application Report.xlsx)", "Browse", 3)

        row4 = ctk.CTkFrame(card, fg_color="transparent")
        row4.grid(row=4, column=0, sticky="ew", padx=16, pady=(12, 16))
        
        _, self.absent_dot, self.absent_lbl = ui.create_status_badge(row4, "Ready")
        self.absent_lbl.master.pack(side="left")

        self.btn_absent = ui.create_primary_button(row4, "Process Data", self.start_process_absent)
        self.btn_absent.pack(side="right")

    # ---------------------------------------------------------
    # Attendance Summary
    # ---------------------------------------------------------
    def _build_att_summary_frame(self):
        f = ctk.CTkFrame(self.main_content, fg_color="transparent")
        self.frames["att_summary"] = f
        f.grid_columnconfigure(0, weight=1)

        ui.create_page_header(f, "Attendance Summary", "Generate employee-wise summary reports from raw attendance portal data.").grid(row=0, column=0, sticky="ew", pady=(0, 20))

        card = ui.create_card(f)
        card.grid(row=1, column=0, sticky="nsew")
        
        ctk.CTkLabel(card, text="Raw Data", font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"), text_color=ui.COLOR_TEXT).grid(row=0, column=0, sticky="w", padx=16, pady=(16, 12))

        ui.create_upload_row(card, self.att_summary_file, "Select Attendance Raw File...", "Browse", 1)

        row2 = ctk.CTkFrame(card, fg_color="transparent")
        row2.grid(row=2, column=0, sticky="ew", padx=16, pady=(12, 16))
        
        _, self.att_dot, self.att_lbl = ui.create_status_badge(row2, "Ready")
        self.att_lbl.master.pack(side="left")

        self.btn_att = ui.create_primary_button(row2, "Summarize Data", self.start_process_att)
        self.btn_att.pack(side="right")

    # ---------------------------------------------------------
    # Time and Leave Master
    # ---------------------------------------------------------
    def _build_time_leave_frame(self):
        f = ctk.CTkFrame(self.main_content, fg_color="transparent")
        self.frames["time_leave"] = f
        f.grid_columnconfigure(0, weight=1)

        hdr = ui.create_page_header(
            f, "Time and Leave Master",
            "Upload single or multi-month records to reconcile Daily Performance Report with Leave and WFH applications."
        )
        hdr.grid(row=0, column=0, sticky="ew", pady=(0, 20))

        card = ui.create_card(f)
        card.grid(row=1, column=0, sticky="nsew")

        ctk.CTkLabel(
            card, text="Data Sources (Single or Multi-Month)",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"),
            text_color=ui.COLOR_TEXT
        ).grid(row=0, column=0, sticky="w", padx=16, pady=(16, 12))

        # 1. Daily performance report
        ui.create_multi_upload_row(
            card, self.tl_perf_files,
            "1. Daily Performance Report (Supports multiple months)",
            placeholder="Select Daily Performance Report Excel/CSV files...",
            row=1
        )

        # 2. Leave application - Active
        ui.create_multi_upload_row(
            card, self.tl_leave_active_files,
            "2. Leave Application - Active (Supports multiple months)",
            placeholder="Select Active Leave Application files...",
            row=2
        )

        # 3. Leave application - Inactive
        ui.create_multi_upload_row(
            card, self.tl_leave_inactive_files,
            "3. Leave Application - Inactive (Supports multiple months)",
            placeholder="Select Inactive Leave Application files...",
            row=3
        )

        # 4. WFH applications
        ui.create_multi_upload_row(
            card, self.tl_wfh_files,
            "4. WFH Applications (Supports multiple months)",
            placeholder="Select WFH Application files...",
            row=4
        )

        # Determinate Progress bar based on actual data volume
        self.tl_prog = ctk.CTkProgressBar(
            card, mode="determinate", height=6, corner_radius=3,
            fg_color=ui.COLOR_INPUT_BG, progress_color=ui.COLOR_ACCENT
        )
        self.tl_prog.grid(row=5, column=0, sticky="ew", padx=16, pady=(8, 4))
        self.tl_prog.set(0)
        self.tl_prog.grid_remove()

        # Real-time progress detail label
        self.tl_prog_lbl = ctk.CTkLabel(
            card, text="",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=11),
            text_color=ui.COLOR_TEXT_SEC, anchor="w", justify="left"
        )
        self.tl_prog_lbl.grid(row=6, column=0, sticky="w", padx=16, pady=(0, 4))
        self.tl_prog_lbl.grid_remove()

        # Action row
        row7 = ctk.CTkFrame(card, fg_color="transparent")
        row7.grid(row=7, column=0, sticky="ew", padx=16, pady=(8, 16))

        _, self.tl_dot, self.tl_lbl = ui.create_status_badge(row7, "Ready")
        self.tl_lbl.master.pack(side="left")

        self.btn_tl = ui.create_primary_button(row7, "Process & Update Report", self.start_process_time_leave, width=170)
        self.btn_tl.pack(side="right")

    # ---------------------------------------------------------
    # Workforce Intelligence Dashboard
    # ---------------------------------------------------------
    def _build_workforce_intelligence_frame(self):
        from workforce_intelligence.dashboard_shell import WorkforceDashboardView
        self.workforce_dashboard_view = WorkforceDashboardView(self.main_content, app=self)
        self.frames["workforce_intelligence"] = self.workforce_dashboard_view

    # ---------------------------------------------------------
    # Time Series Analysis
    # ---------------------------------------------------------
    def _build_time_series_frame(self):
        f = ctk.CTkScrollableFrame(self.main_content, fg_color="transparent")
        self.frames["analyse_time_series"] = f
        self.frames["analyse_dashboard"] = f  # backwards-compatible alias
        f.grid_columnconfigure(0, weight=1)

        # Page Header
        hdr = ui.create_page_header(
            f, "Time Series Analysis",
            "Workforce compliance, exceptions, biometric swipes, and multi-level breakdown."
        )
        hdr.grid(row=0, column=0, sticky="ew", pady=(0, 16))

        # 1. Dataset Status Header Card
        self.ts_banner_card = ui.create_card(f)
        self.ts_banner_card.grid(row=1, column=0, sticky="ew", pady=(0, 14))
        self.ts_banner_card.grid_columnconfigure(0, weight=1)

        b_frame = ctk.CTkFrame(self.ts_banner_card, fg_color="transparent")
        b_frame.grid(row=0, column=0, sticky="ew", padx=16, pady=12)
        b_frame.grid_columnconfigure(0, weight=1)

        self.ts_banner_lbl = ctk.CTkLabel(
            b_frame, text="⚠️ No dataset loaded yet. Please upload a dataset in the Upload section to begin analysis.",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=13),
            text_color=ui.COLOR_WARNING, anchor="w"
        )
        self.ts_banner_lbl.grid(row=0, column=0, sticky="w")

        self.btn_ts_go_upload = ui.create_secondary_button(
            b_frame, "📂 Go to Upload Section",
            lambda: self.select_frame_by_name("analyse_upload"), width=180
        )
        self.btn_ts_go_upload.grid(row=0, column=1, sticky="e", padx=(10, 0))

        # 2. Filter & Actions Toolbar
        self.card_filters = ui.create_card(f)
        self.card_filters.grid(row=2, column=0, sticky="ew", pady=(0, 14))

        f_header = ctk.CTkFrame(self.card_filters, fg_color="transparent")
        f_header.grid(row=0, column=0, sticky="ew", padx=16, pady=(12, 6))
        ctk.CTkLabel(
            f_header, text="Filters & Scope",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=14, weight="bold"),
            text_color=ui.COLOR_TEXT
        ).pack(side="left")

        self.btn_ts_export = ui.create_secondary_button(
            f_header, "📥 Export Report to Excel", self._export_ts_excel, width=180
        )
        self.btn_ts_export.pack(side="right")

        f_body = ctk.CTkFrame(self.card_filters, fg_color="transparent")
        f_body.grid(row=1, column=0, sticky="ew", padx=16, pady=(0, 12))
        f_body.grid_columnconfigure((0, 1, 2, 3), weight=1)

        # BU Combobox
        ctk.CTkLabel(f_body, text="Business Unit", font=ctk.CTkFont(size=11, weight="bold"), text_color=ui.COLOR_TEXT_DIM).grid(row=0, column=0, sticky="w", padx=4, pady=(0, 2))
        self.ts_bu_combo = ctk.CTkComboBox(
            f_body, variable=self.ts_bu_var, values=["All Business Units"],
            command=self._on_ts_filter_changed, height=32,
            fg_color=ui.COLOR_INPUT_BG, border_color=ui.COLOR_BORDER,
            text_color=ui.COLOR_TEXT, dropdown_fg_color=ui.COLOR_CARD,
        )
        self.ts_bu_combo.grid(row=1, column=0, sticky="ew", padx=4)

        # Department Combobox
        ctk.CTkLabel(f_body, text="Department", font=ctk.CTkFont(size=11, weight="bold"), text_color=ui.COLOR_TEXT_DIM).grid(row=0, column=1, sticky="w", padx=4, pady=(0, 2))
        self.ts_dept_combo = ctk.CTkComboBox(
            f_body, variable=self.ts_dept_var, values=["All Departments"],
            command=self._on_ts_filter_changed, height=32,
            fg_color=ui.COLOR_INPUT_BG, border_color=ui.COLOR_BORDER,
            text_color=ui.COLOR_TEXT, dropdown_fg_color=ui.COLOR_CARD,
        )
        self.ts_dept_combo.grid(row=1, column=1, sticky="ew", padx=4)

        # Month Combobox
        ctk.CTkLabel(f_body, text="Month", font=ctk.CTkFont(size=11, weight="bold"), text_color=ui.COLOR_TEXT_DIM).grid(row=0, column=2, sticky="w", padx=4, pady=(0, 2))
        self.ts_month_combo = ctk.CTkComboBox(
            f_body, variable=self.ts_month_var, values=["All Months"],
            command=self._on_ts_filter_changed, height=32,
            fg_color=ui.COLOR_INPUT_BG, border_color=ui.COLOR_BORDER,
            text_color=ui.COLOR_TEXT, dropdown_fg_color=ui.COLOR_CARD,
        )
        self.ts_month_combo.grid(row=1, column=2, sticky="ew", padx=4)

        # Reset button
        btn_reset = ui.create_secondary_button(f_body, "Reset Filters", self._reset_ts_filters, width=100)
        btn_reset.grid(row=1, column=3, sticky="ew", padx=4)

        # 3. KPI Metrics Grid (8 cards, 4 per row)
        self.ts_kpi_frame = ctk.CTkFrame(f, fg_color="transparent")
        self.ts_kpi_frame.grid(row=3, column=0, sticky="ew", pady=(0, 14))
        self.ts_kpi_frame.grid_columnconfigure((0, 1, 2, 3), weight=1, uniform="kpi")

        # Row 0
        c1, self.lbl_kpi_leave_val, self.lbl_kpi_leave_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "Leave Compliance Rate")
        c1.grid(row=0, column=0, sticky="nsew", padx=4, pady=4)

        c2, self.lbl_kpi_appr_val, self.lbl_kpi_appr_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "Manager Approval Compliance")
        c2.grid(row=0, column=1, sticky="nsew", padx=4, pady=4)

        c3, self.lbl_kpi_reg_val, self.lbl_kpi_reg_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "Attendance Exception Rate")
        c3.grid(row=0, column=2, sticky="nsew", padx=4, pady=4)

        c4, self.lbl_kpi_wfh_val, self.lbl_kpi_wfh_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "WFH Exception Rate (>3d)")
        c4.grid(row=0, column=3, sticky="nsew", padx=4, pady=4)

        # Row 1
        c5, self.lbl_kpi_rep_val, self.lbl_kpi_rep_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "Repeat Attendance Non-Compliance")
        c5.grid(row=1, column=0, sticky="nsew", padx=4, pady=4)

        c6, self.lbl_kpi_in_val, self.lbl_kpi_in_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "AVG In Time (Present/MS)")
        c6.grid(row=1, column=1, sticky="nsew", padx=4, pady=4)

        c7, self.lbl_kpi_out_val, self.lbl_kpi_out_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "AVG Out Time (Present/MS)")
        c7.grid(row=1, column=2, sticky="nsew", padx=4, pady=4)

        c8, self.lbl_kpi_hrs_val, self.lbl_kpi_hrs_sub = ui.create_kpi_metric_card(self.ts_kpi_frame, "AVG Working Hours")
        c8.grid(row=1, column=3, sticky="nsew", padx=4, pady=4)

        # 4. Multi-Level Breakdown Table Card
        card_table = ui.create_card(f)
        card_table.grid(row=4, column=0, sticky="ew", pady=(0, 20))

        tbl_top = ctk.CTkFrame(card_table, fg_color="transparent")
        tbl_top.grid(row=0, column=0, sticky="ew", padx=16, pady=(12, 8))

        ctk.CTkLabel(
            tbl_top, text="Multi-Level Breakdown Analysis",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"),
            text_color=ui.COLOR_TEXT
        ).pack(side="left", padx=(0, 16))

        # Segmented Control for Levels
        self.ts_level_seg = ctk.CTkSegmentedButton(
            tbl_top,
            values=["Business Unit", "Department", "Reporting Manager", "Employee"],
            command=self._on_ts_level_changed,
            selected_color=ui.COLOR_ACCENT,
            selected_hover_color=ui.COLOR_ACCENT_HOVER,
            unselected_color=ui.COLOR_INPUT_BG,
            unselected_hover_color=ui.COLOR_BTN_SEC_HOV,
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=12, weight="bold"),
        )
        self.ts_level_seg.set("Business Unit")
        self.ts_level_seg.pack(side="left")

        # Search box on right
        search_box = ctk.CTkEntry(
            tbl_top,
            textvariable=self.ts_search_var,
            placeholder_text="Search entity name...",
            width=200, height=32,
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=12),
            fg_color=ui.COLOR_INPUT_BG, border_color=ui.COLOR_BORDER,
            border_width=1, text_color=ui.COLOR_TEXT,
        )
        search_box.pack(side="right")
        self.ts_search_var.trace_add("write", self._on_ts_search_changed)

        # Treeview container
        tree_container = ctk.CTkFrame(card_table, fg_color=ui.COLOR_CARD)
        tree_container.grid(row=1, column=0, sticky="ew", padx=16, pady=(0, 14))
        tree_container.grid_columnconfigure(0, weight=1)

        # Style Treeview
        style = ttk.Style()
        style.theme_use("clam")
        style.configure(
            "TS.Treeview",
            background=ui.COLOR_CARD,
            foreground=ui.COLOR_TEXT,
            fieldbackground=ui.COLOR_CARD,
            rowheight=28,
            font=(ui.FONT_FAMILY, 10),
            borderwidth=0,
        )
        style.configure(
            "TS.Treeview.Heading",
            background="#1A2340",
            foreground="#00FF99",
            font=(ui.FONT_FAMILY, 10, "bold"),
            borderwidth=0,
            relief="flat",
        )
        style.map(
            "TS.Treeview",
            background=[("selected", ui.COLOR_ACCENT)],
            foreground=[("selected", "#FFFFFF")],
        )
        style.map(
            "TS.Treeview.Heading",
            background=[("active", "#243052")],
        )

        cols = (
            "entity", "records", "emps", "leave_days", "leave_emp_pct",
            "appr_days", "appr_mgr_pct", "regularized", "wfh_excess",
            "repeat_dev", "in_time", "out_time", "work_hrs"
        )
        self.ts_tree = ttk.Treeview(
            tree_container, columns=cols, show="headings",
            style="TS.Treeview", height=10
        )

        col_configs = [
            ("entity", "Entity Name", 160, "w"),
            ("records", "Rows", 65, "center"),
            ("emps", "Emps", 60, "center"),
            ("leave_days", "Avg Leave Apply", 110, "center"),
            ("leave_emp_pct", "Emp Leave %", 90, "center"),
            ("appr_days", "Avg Approval", 95, "center"),
            ("appr_mgr_pct", "Mgr Appr %", 85, "center"),
            ("regularized", "Regularized", 85, "center"),
            ("wfh_excess", "WFH >3 Days", 85, "center"),
            ("repeat_dev", "Repeat Dev", 80, "center"),
            ("in_time", "AVG In Time", 95, "center"),
            ("out_time", "AVG Out Time", 95, "center"),
            ("work_hrs", "AVG Work Hrs", 95, "center"),
        ]

        for col_id, col_text, col_w, col_anchor in col_configs:
            self.ts_tree.heading(col_id, text=col_text)
            self.ts_tree.column(col_id, width=col_w, minwidth=60, anchor=col_anchor)

        v_scroll = ttk.Scrollbar(tree_container, orient="vertical", command=self.ts_tree.yview)
        h_scroll = ttk.Scrollbar(tree_container, orient="horizontal", command=self.ts_tree.xview)
        self.ts_tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        self.ts_tree.grid(row=0, column=0, sticky="ew")
        v_scroll.grid(row=0, column=1, sticky="ns")
        h_scroll.grid(row=1, column=0, sticky="ew")

    # ---------------------------------------------------------
    # Analyse Upload
    # ---------------------------------------------------------
    def _build_analyse_upload_frame(self):
        f = ctk.CTkFrame(self.main_content, fg_color="transparent")
        self.frames["analyse_upload"] = f
        f.grid_columnconfigure(0, weight=1)
        f.grid_rowconfigure(2, weight=1)

        hdr = ui.create_page_header(
            f, "Upload Dataset",
            "Upload Daily Performance report or attendance logs (.xlsx, .xls, .csv) for Time Series Analysis."
        )
        hdr.grid(row=0, column=0, sticky="ew", pady=(0, 14))

        actions = ctk.CTkFrame(hdr, fg_color="transparent")
        actions.grid(row=0, column=1, sticky="e")
        ui.create_secondary_button(actions, "← Workforce Intelligence", lambda: self.select_frame_by_name("workforce_intelligence"), 180).pack(side="left", padx=(0, 10))
        ui.create_secondary_button(actions, "← Time Series Analysis", lambda: self.select_frame_by_name("analyse_time_series"), 180).pack(side="left")

        # Upload Card
        card = ui.create_card(f)
        card.grid(row=1, column=0, sticky="nsew", pady=(0, 12))

        ctk.CTkLabel(
            card, text="1. Select Attendance / Performance Dataset",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"),
            text_color=ui.COLOR_TEXT
        ).grid(row=0, column=0, sticky="w", padx=16, pady=(12, 2))

        ctk.CTkLabel(
            card, text="Supported formats: Excel (.xlsx, .xls) and CSV (.csv).",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=12),
            text_color=ui.COLOR_TEXT_SEC
        ).grid(row=1, column=0, sticky="w", padx=16, pady=(0, 8))

        ui.create_upload_row(card, self.ts_file_path, "Select performance dataset to ingest...", "Browse", 2)

        self.ts_prog = ctk.CTkProgressBar(
            card, mode="indeterminate", height=4, corner_radius=2,
            fg_color=ui.COLOR_INPUT_BG, progress_color=ui.COLOR_ACCENT
        )
        self.ts_prog.grid(row=3, column=0, sticky="ew", padx=16, pady=(4, 0))
        self.ts_prog.set(0)
        self.ts_prog.grid_remove()

        row4 = ctk.CTkFrame(card, fg_color="transparent")
        row4.grid(row=4, column=0, sticky="ew", padx=16, pady=(8, 12))

        _, self.ts_dot, self.ts_lbl = ui.create_status_badge(row4, "Ready")
        self.ts_lbl.master.pack(side="left")

        self.btn_ts_load = ui.create_primary_button(row4, "Load & Ingest Dataset", self.start_ts_analysis, width=170)
        self.btn_ts_load.pack(side="right")

        finbox_sample = Path("local_data/Daily Performance Report 01 Sep 2026 - 13 Sep 2026 - FinBox.xlsx")
        if finbox_sample.exists():
            def _load_sample():
                self.ts_file_path.set(str(finbox_sample.resolve()))
                self.start_ts_analysis()
            ui.create_secondary_button(row4, "📂 Load FinBox Sample", _load_sample, width=170).pack(side="right", padx=(0, 10))

        # Dataset Ingestion Summary Card
        res_card = ui.create_card(f)
        res_card.grid(row=2, column=0, sticky="nsew")
        res_card.grid_columnconfigure(0, weight=1)
        res_card.grid_rowconfigure(1, weight=1)

        res_top = ctk.CTkFrame(res_card, fg_color="transparent")
        res_top.grid(row=0, column=0, sticky="ew", padx=16, pady=(12, 6))

        ctk.CTkLabel(
            res_top, text="Dataset Summary & Metadata",
            font=ctk.CTkFont(family=ui.FONT_FAMILY, size=15, weight="bold"),
            text_color=ui.COLOR_TEXT
        ).pack(side="left")

        self.btn_ts_open_tsa = ui.create_primary_button(
            res_top, "🚀 Open Time Series Analysis →",
            lambda: self.select_frame_by_name("analyse_time_series"), width=230
        )
        self.btn_ts_open_tsa.pack(side="right")
        self.btn_ts_open_tsa.configure(state="disabled")

        self.ts_upload_summary_text = ctk.CTkTextbox(
            res_card, height=140, font=ctk.CTkFont(family="Consolas", size=12),
            fg_color=ui.COLOR_INPUT_BG, border_color=ui.COLOR_BORDER,
            border_width=1, text_color=ui.COLOR_TEXT, corner_radius=6
        )
        self.ts_upload_summary_text.grid(row=1, column=0, sticky="nsew", padx=16, pady=(0, 12))
        self.ts_upload_summary_text.insert("0.0", "Select an attendance or performance dataset above and click 'Load & Ingest Dataset'.\nOnce loaded, all metrics and drilldowns will be instantly populated in Time Series Analysis.")
        self.ts_upload_summary_text.configure(state="disabled")



    # =========================================================
    # Time Series Analysis Event Handlers
    # =========================================================

    def start_ts_analysis(self):
        file_path = self.ts_file_path.get().strip()
        if not file_path:
            messagebox.showerror("Error", "Please select a dataset file to analyse.")
            return
        if not os.path.exists(file_path):
            messagebox.showerror("Error", f"File not found:\n{file_path}")
            return

        self.is_ts_loading = True
        self.btn_ts_load.configure(state="disabled")
        self.ts_prog.grid()
        self.ts_prog.start()
        ui.update_status(self.ts_dot, self.ts_lbl, "Loading and indexing dataset...", "processing")

        self._latest_ts_job_id += 1
        job_id = self._latest_ts_job_id
        threading.Thread(target=self._run_ts_load_job, args=(file_path, job_id), daemon=True).start()

    def _run_ts_load_job(self, file_path, job_id):
        try:
            snapshot = snapshot_service.prepare_dataset(file_path)
            self.after(0, lambda s=snapshot, p=file_path, j=job_id: self._ts_load_success(s, p, j))
        except Exception as e:
            err_msg = str(e)
            self.after(0, lambda m=err_msg, j=job_id: self._ts_load_error(m, j))

    def _ts_load_success(self, snapshot, file_path, job_id):
        if job_id < self._latest_ts_job_id:
            return  # Ignore stale completion

        self.is_ts_loading = False
        self.btn_ts_load.configure(state="normal")
        self.ts_prog.stop()
        self.ts_prog.grid_remove()

        self.ts_dataset = snapshot.fact_df
        total_rows = snapshot.row_count
        total_emps = snapshot.metadata.get("employee_count", 0)
        p_name = Path(file_path).name

        ui.update_status(self.ts_dot, self.ts_lbl, f"Active: {total_rows:,} records • {total_emps:,} employees", "success")

        # Update Workforce Intelligence Dashboard header metadata
        if hasattr(self, "workforce_dashboard_view"):
            self.workforce_dashboard_view.sync_snapshot_state(snapshot)

        # Enable Open Time Series Analysis button in Upload section
        if hasattr(self, "btn_ts_open_tsa"):
            self.btn_ts_open_tsa.configure(state="normal")

        # Update Summary text in Upload section
        if hasattr(self, "ts_upload_summary_text"):
            self.ts_upload_summary_text.configure(state="normal")
            self.ts_upload_summary_text.delete("0.0", "end")
            bus = snapshot.filter_options["business_units"]
            depts = snapshot.filter_options["departments"]
            rms = snapshot.filter_options["managers"]
            months = snapshot.filter_options["months"]
            min_d = snapshot.metadata.get("min_date", "N/A")
            max_d = snapshot.metadata.get("max_date", "N/A")

            summary_info = (
                f"✓ DATASET INGESTED SUCCESSFULLY\n"
                f"{'=' * 60}\n"
                f"File Name:            {p_name}\n"
                f"Total Row Count:      {total_rows:,}\n"
                f"Total Unique Emps:    {total_emps:,}\n"
                f"Date Range:           {min_d} to {max_d}\n"
                f"Business Units ({len(bus)}):   {', '.join(bus[:8])}{'...' if len(bus) > 8 else ''}\n"
                f"Departments ({len(depts)}):      {', '.join(depts[:8])}{'...' if len(depts) > 8 else ''}\n"
                f"Managers ({len(rms)}):         {', '.join(rms[:8])}{'...' if len(rms) > 8 else ''}\n"
                f"Months ({len(months)}):           {', '.join(months)}\n"
                f"{'=' * 60}\n"
                f"Status: Ready for Time Series Analysis! Click 'Open Time Series Analysis →' above."
            )
            self.ts_upload_summary_text.insert("0.0", summary_info)
            self.ts_upload_summary_text.configure(state="disabled")

        # Update Time Series banner
        if hasattr(self, "ts_banner_lbl"):
            self.ts_banner_lbl.configure(
                text=f"📊 Active Dataset: {p_name}  •  {total_rows:,} records  •  {total_emps:,} employees",
                text_color=ui.COLOR_SUCCESS
            )
            self.btn_ts_go_upload.configure(text="📂 Change Dataset")

        # Populate filters
        self.ts_bu_combo.configure(values=["All Business Units"] + bus)
        self.ts_bu_var.set("All Business Units")

        self.ts_dept_combo.configure(values=["All Departments"] + depts)
        self.ts_dept_var.set("All Departments")

        self.ts_month_combo.configure(values=["All Months"] + months)
        self.ts_month_var.set("All Months")

        self._refresh_ts_dashboard()

    def _ts_load_error(self, err_msg, job_id):
        if job_id < self._latest_ts_job_id:
            return  # Ignore stale error

        self.is_ts_loading = False
        self.btn_ts_load.configure(state="normal")
        self.ts_prog.stop()
        self.ts_prog.grid_remove()

        active_snap = snapshot_service.get_active_snapshot()
        if active_snap and active_snap.is_valid():
            total_rows = active_snap.row_count
            total_emps = active_snap.metadata.get("employee_count", 0)
            ui.update_status(self.ts_dot, self.ts_lbl, f"Active: {total_rows:,} records • {total_emps:,} employees", "success")
            if active_snap.raw_source:
                self.ts_file_path.set(str(active_snap.raw_source))
        else:
            ui.update_status(self.ts_dot, self.ts_lbl, "Failed to load dataset", "error")

        messagebox.showerror("Dataset Loading Error", f"Could not load and process dataset:\n{err_msg}")

    def _on_ts_filter_changed(self, choice=None):
        self._refresh_ts_dashboard()

    def _on_ts_level_changed(self, choice):
        self.ts_level_var.set(choice)
        self._refresh_ts_table(reload_data=True)

    def _reset_ts_filters(self):
        self.ts_bu_var.set("All Business Units")
        self.ts_dept_var.set("All Departments")
        self.ts_month_var.set("All Months")
        self.ts_search_var.set("")
        self._refresh_ts_dashboard()

    def _refresh_ts_dashboard(self):
        if not snapshot_service.has_active_snapshot():
            return

        bu = self.ts_bu_var.get()
        dept = self.ts_dept_var.get()
        month = self.ts_month_var.get()

        metrics = snapshot_service.get_dashboard_metrics(
            business_unit=bu, department=dept, month=month
        )

        # 1. Leave Compliance
        leave_days = metrics['avg_leave_apply_days']
        if leave_days >= 0:
            self.lbl_kpi_leave_val.configure(text=f"{leave_days:.1f} days")
            self.lbl_kpi_leave_sub.configure(text=f"In Advance • Emp: {metrics['applied_by_emp_pct']}% | Admin: {metrics['applied_by_admin_pct']}%")
        else:
            self.lbl_kpi_leave_val.configure(text=f"{abs(leave_days):.1f} days late")
            self.lbl_kpi_leave_sub.configure(text=f"Post-Leave • Emp: {metrics['applied_by_emp_pct']}% | Admin: {metrics['applied_by_admin_pct']}%")

        # 2. Manager Approval Compliance
        self.lbl_kpi_appr_val.configure(text=f"{metrics['avg_approval_days']:.1f} days")
        self.lbl_kpi_appr_sub.configure(text=f"Manager: {metrics['appr_mgr_pct']}%  |  Admin: {metrics['appr_admin_pct']}%")

        # 3. Attendance Exceptions
        self.lbl_kpi_reg_val.configure(text=f"{metrics['regularized_days']:,} days")
        self.lbl_kpi_reg_sub.configure(text=f"Regularization Rate: {metrics['reg_rate_pct']}%")

        # 4. WFH Exceptions (>3 Days)
        self.lbl_kpi_wfh_val.configure(text=f"{metrics['wfh_excess_days']:,} days")
        self.lbl_kpi_wfh_sub.configure(text=f"{metrics['wfh_violating_emp_count']} emps > 3 days ({metrics['wfh_exception_rate_pct']}%)")

        # 5. Repeat Non-Compliance
        self.lbl_kpi_rep_val.configure(text=f"{metrics['repeat_emp_count']:,} Emps ({metrics['repeat_emp_pct']}%)")
        self.lbl_kpi_rep_sub.configure(text=f"{metrics['repeat_rm_count']} Managers with repeat deviations")

        # 6. AVG In Time
        self.lbl_kpi_in_val.configure(text=metrics['avg_in_time'])
        self.lbl_kpi_in_sub.configure(text="Present & Missing Swipes")

        # 7. AVG Out Time
        self.lbl_kpi_out_val.configure(text=metrics['avg_out_time'])
        self.lbl_kpi_out_sub.configure(text="Present & Missing Swipes")

        # 8. AVG Working Hours
        self.lbl_kpi_hrs_val.configure(text=metrics['avg_working_hours'])
        self.lbl_kpi_hrs_sub.configure(text=f"{metrics['avg_working_hours_decimal']} decimal hours (P & MS)")

        # Refresh table
        self._refresh_ts_table(reload_data=True)

    def _refresh_ts_table(self, reload_data=True):
        if not snapshot_service.has_active_snapshot():
            return

        level = self.ts_level_var.get()
        bu = self.ts_bu_var.get()
        dept = self.ts_dept_var.get()
        month = self.ts_month_var.get()

        if reload_data or not hasattr(self, "_current_ts_breakdown_rows") or not self._current_ts_breakdown_rows:
            self._current_ts_breakdown_rows = snapshot_service.get_breakdown(
                level=level,
                business_unit=bu,
                department=dept,
                month=month,
            )

        self._render_ts_table_rows()

    def _on_ts_search_changed(self, *args):
        if hasattr(self, "_search_debounce_id") and self._search_debounce_id:
            try:
                self.after_cancel(self._search_debounce_id)
            except Exception:
                pass
        self._search_debounce_id = self.after(100, self._render_ts_table_rows)

    def _render_ts_table_rows(self):
        for item in self.ts_tree.get_children():
            self.ts_tree.delete(item)

        search_q = self.ts_search_var.get().strip().lower()
        rows = getattr(self, "_current_ts_breakdown_rows", [])

        for r in rows:
            e_name = str(r.get("Entity Name", ""))
            e_id = str(r.get("Entity ID", ""))
            if search_q and (search_q not in e_name.lower() and search_q not in e_id.lower()):
                continue

            vals = (
                e_name if e_name else e_id,
                f"{r.get('Total Records', 0):,}",
                f"{r.get('Unique Employees', 0):,}",
                f"{r.get('Avg Days to Apply Leave', 0.0)}d",
                r.get("Leave Applied % (Emp)", "0.0%"),
                f"{r.get('Avg Approval Days', 0.0)}d",
                r.get("Approval % (Mgr)", "0.0%"),
                f"{r.get('Regularized Days', 0):,}",
                f"{r.get('WFH Exception Days (>3)', 0):,}",
                f"{r.get('Repeat Deviant Emps', 0):,}",
                r.get("Avg In Time", "-"),
                r.get("Avg Out Time", "-"),
                r.get("Avg Working Hours", "-"),
            )
            self.ts_tree.insert("", "end", values=vals)

    def _export_ts_excel(self):
        if not snapshot_service.has_active_snapshot():
            messagebox.showwarning("Warning", "Please load a dataset first before exporting.")
            return

        out_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"Time_Series_Analysis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not out_path:
            return

        try:
            bu = self.ts_bu_var.get()
            dept = self.ts_dept_var.get()
            month = self.ts_month_var.get()
            saved_p = snapshot_service.export_report(
                output_path=out_path,
                business_unit=bu,
                department=dept,
                month=month,
            )
            messagebox.showinfo(
                "Export Complete",
                f"Time Series Analysis report successfully exported to:\n{saved_p}"
            )
        except Exception as e:
            messagebox.showerror("Export Failed", f"Failed to export Excel report:\n{e}")

    # =========================================================
    # Event Handlers (UI updates)
    # =========================================================

    def start_analyse(self):
        file_path = self.analyse_upload_file.get().strip()
        if not file_path:
            messagebox.showerror("Error", "Please select a file to analyse.")
            return
        if not os.path.exists(file_path):
            messagebox.showerror("Error", f"File not found:\n{file_path}")
            return

        self.is_analysing = True
        self.btn_analyse.configure(state="disabled")
        self.analyse_prog.grid()
        self.analyse_prog.start()
        ui.update_status(self.analyse_dot, self.analyse_lbl, "Analysing dataset...", "processing")

        threading.Thread(target=self._run_analyse_job, args=(file_path,), daemon=True).start()

    def _run_analyse_job(self, file_path):
        try:
            p = Path(file_path)
            ext = p.suffix.lower()
            file_size_kb = p.stat().st_size / 1024

            lines = []
            lines.append("=" * 64)
            lines.append(" DATASET ANALYSIS REPORT")
            lines.append("=" * 64)
            lines.append(f"File:      {p.name}")
            lines.append(f"Path:      {file_path}")
            lines.append(f"Size:      {file_size_kb:.1f} KB")
            lines.append(f"Analyzed:  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
            lines.append("-" * 64)

            if ext == ".csv":
                df = pd.read_csv(file_path)
                lines.append("Format:    CSV (Single Table)")
                lines.append(f"Total Rows:    {len(df):,}")
                lines.append(f"Total Columns: {len(df.columns):,}")
                lines.append(f"Columns:       {', '.join(df.columns.astype(str))}")
                lines.append("")
                lines.append("Column Details:")
                for col in df.columns:
                    non_null = df[col].count()
                    null_cnt = df[col].isna().sum()
                    dtype = df[col].dtype
                    lines.append(f"  • {col}: {dtype} | {non_null:,} non-null, {null_cnt:,} nulls")
            else:
                with pd.ExcelFile(file_path) as xl:
                    lines.append("Format:    Excel Workbook")
                    lines.append(f"Sheets ({len(xl.sheet_names)}): {', '.join(xl.sheet_names)}")
                    lines.append("")
                    for sheet in xl.sheet_names:
                        df = xl.parse(sheet)
                        lines.append(f"--- Sheet: '{sheet}' ---")
                        lines.append(f"  Rows:    {len(df):,}")
                        lines.append(f"  Columns: {len(df.columns):,}")
                        cols_preview = ', '.join(df.columns.astype(str)[:10])
                        if len(df.columns) > 10:
                            cols_preview += "..."
                        lines.append(f"  Columns: {cols_preview}")
                        null_sum = df.isna().sum().sum()
                        lines.append(f"  Total Missing Cells: {null_sum:,}")
                        lines.append("")

            lines.append("=" * 64)
            lines.append("Status: Analysis complete.")
            report = "\n".join(lines)
            self.after(0, lambda: self._analyse_success(report))
        except Exception as e:
            self.after(0, lambda: self._analyse_error(str(e)))

    def _analyse_success(self, report):
        self.is_analysing = False
        self.btn_analyse.configure(state="normal")
        self.analyse_prog.stop()
        self.analyse_prog.grid_remove()
        ui.update_status(self.analyse_dot, self.analyse_lbl, "Analysis Complete", "success")

        self.analyse_result_text.configure(state="normal")
        self.analyse_result_text.delete("0.0", "end")
        self.analyse_result_text.insert("0.0", report)
        self.analyse_result_text.configure(state="disabled")

    def _analyse_error(self, err):
        self.is_analysing = False
        self.btn_analyse.configure(state="normal")
        self.analyse_prog.stop()
        self.analyse_prog.grid_remove()
        ui.update_status(self.analyse_dot, self.analyse_lbl, "Analysis Failed", "error")
        messagebox.showerror("Analysis Error", f"Failed to analyze file:\n{err}")


    def start_transform(self):
        if not self.selected_input_file.get():
            messagebox.showerror("Error", "Please select a file.")
            return
        self.is_processing = True
        self.btn_kra_tf.configure(state="disabled")
        self.kra_prog.grid()
        self.kra_prog.start()
        ui.update_status(self.kra_dot, self.kra_lbl, "Transforming...", "processing")
        threading.Thread(target=self._run_transform_job, daemon=True).start()

    def cancel_transform(self): pass # Unused currently, logic can be complex

    def _transform_success(self, out_path, cnt):
        self.is_processing = False
        self.btn_kra_tf.configure(state="normal")
        self.kra_prog.stop()
        self.kra_prog.grid_remove()
        ui.update_status(self.kra_dot, self.kra_lbl, f"Success • {cnt} sheets", "success")
        messagebox.showinfo("Done", f"File transformed successfully.\n\nSaved to:\n{out_path}")

    def _transform_error(self, err):
        self.is_processing = False
        self.btn_kra_tf.configure(state="normal")
        self.kra_prog.stop()
        self.kra_prog.grid_remove()
        ui.update_status(self.kra_dot, self.kra_lbl, "Failed", "error")
        messagebox.showerror("Error", err)

    def start_generate(self):
        mode = self.gen_mode_seg.get() if hasattr(self, 'gen_mode_seg') else "Existing KRA Upload"
        file_path = self.selected_okr_input_file.get() if mode == "OKR Upload" else self.selected_gen_input_file.get()
        if not file_path:
            messagebox.showerror("Error", "Please select a file.")
            return
        self.is_generating = True
        self.btn_gen.configure(state="disabled")
        self.gen_prog.grid()
        self.gen_prog.start()
        ui.update_status(self.gen_dot, self.gen_lbl, "Generating...", "processing")
        threading.Thread(target=self._run_generate_job, daemon=True).start()

    def cancel_generate(self): pass

    def _reset_gen_ui_state(self):
        self.is_generating = False

        self.btn_gen.configure(state="normal")
        self.gen_prog.stop()
        self.gen_prog.grid_remove()

    def _generate_success(self, out_path):
        self._reset_gen_ui_state()
        ui.update_status(self.gen_dot, self.gen_lbl, "Success", "success")
        messagebox.showinfo("Done", f"Generated successfully:\n{out_path}")

    def _generate_error(self, err):
        self._reset_gen_ui_state()
        ui.update_status(self.gen_dot, self.gen_lbl, "Failed", "error")
        messagebox.showerror("Error", err)

    def start_process_absent(self):
        if not (self.absent_emp_file.get() and self.absent_att_file.get() and self.absent_wfh_file.get()):
            messagebox.showerror("Error", "Please select all 3 files.")
            return
        self.btn_absent.configure(state="disabled")
        ui.update_status(self.absent_dot, self.absent_lbl, "Processing...", "processing")
        threading.Thread(target=self._run_absent_job, daemon=True).start()

    def _absent_success(self, path):
        self.btn_absent.configure(state="normal")
        ui.update_status(self.absent_dot, self.absent_lbl, "Success", "success")
        messagebox.showinfo("Done", f"Saved to:\n{path}")

    def _absent_error(self, err):
        self.btn_absent.configure(state="normal")
        ui.update_status(self.absent_dot, self.absent_lbl, "Failed", "error")
        messagebox.showerror("Error", err)

    def start_process_att(self):
        if not self.att_summary_file.get():
            messagebox.showerror("Error", "Select file first.")
            return
        self.btn_att.configure(state="disabled")
        ui.update_status(self.att_dot, self.att_lbl, "Processing...", "processing")
        threading.Thread(target=self._run_att_job, daemon=True).start()

    def _att_success(self, path):
        self.btn_att.configure(state="normal")
        ui.update_status(self.att_dot, self.att_lbl, "Success", "success")
        messagebox.showinfo("Done", f"Saved to:\n{path}")

    def _att_error(self, err):
        self.btn_att.configure(state="normal")
        ui.update_status(self.att_dot, self.att_lbl, "Failed", "error")
        messagebox.showerror("Error", err)

    def start_process_time_leave(self):
        if not self.tl_perf_files:
            messagebox.showerror("Error", "Please select at least one Daily Performance Report file.")
            return
        if not (self.tl_leave_active_files or self.tl_leave_inactive_files or self.tl_wfh_files):
            if not messagebox.askyesno("Confirm", "No Leave or WFH application files were selected. Proceed anyway?"):
                return

        out_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=f"Daily_Performance_Report_Updated_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
        if not out_path:
            return

        self.btn_tl.configure(state="disabled")
        self.tl_prog.grid()
        self.tl_prog.set(0)
        self.tl_prog_lbl.grid()
        self.tl_prog_lbl.configure(text="Initializing data reconciliation...")
        ui.update_status(self.tl_dot, self.tl_lbl, "Starting (0%)...", "processing")
        threading.Thread(target=self._run_time_leave_job, args=(out_path,), daemon=True).start()

    def _update_tl_progress(self, pct: float, msg: str):
        self.tl_prog.set(max(0.0, min(1.0, pct)))
        self.tl_prog_lbl.configure(text=msg)
        pct_int = int(round(pct * 100))
        ui.update_status(self.tl_dot, self.tl_lbl, f"Processing ({pct_int}%)", "processing")

    def _run_time_leave_job(self, out_path):
        def _on_progress(pct: float, msg: str):
            self.after(0, lambda p=pct, m=msg: self._update_tl_progress(p, m))

        try:
            import time_leave_master
            import importlib
            importlib.reload(time_leave_master)
            stats = time_leave_master.reconcile_time_and_leave(
                perf_files=self.tl_perf_files,
                leave_active_files=self.tl_leave_active_files,
                leave_inactive_files=self.tl_leave_inactive_files,
                wfh_files=self.tl_wfh_files,
                output_path=out_path,
                progress_callback=_on_progress
            )
            self.after(0, lambda s=stats, o=out_path: self._time_leave_success(s, o))
        except Exception as e:
            err_msg = str(e)
            self.after(0, lambda m=err_msg: self._time_leave_error(m))


    def _time_leave_success(self, stats, out_path):
        self.btn_tl.configure(state="normal")
        self.tl_prog.set(1.0)
        self.tl_prog_lbl.configure(text=f"Completed {stats['total_rows']:,} rows • {stats['matched_leave_count']:,} leaves, {stats['matched_wfh_count']:,} WFH matched")
        ui.update_status(self.tl_dot, self.tl_lbl, "Success", "success")

        msg = (
            f"Daily Performance Report updated successfully!\n\n"
            f"• Total Rows: {stats['total_rows']:,}\n"
            f"• Matched Leave records: {stats['matched_leave_count']:,}\n"
            f"• Matched WFH records: {stats['matched_wfh_count']:,}\n"
        )
        if stats.get('quantity_violation_count', 0) > 0:
            msg += f"\n⚠️ Warning: {stats['quantity_violation_count']} date entries have total quantity > 1.0 per employee."
        if stats.get('unmatched_applications_count', 0) > 0:
            msg += f"\nℹ️ Unmatched Applications: {stats['unmatched_applications_count']:,} applications could not be matched."

        msg += f"\n\nMain Output Saved to:\n{out_path}"
        if stats.get('error_log_path'):
            msg += f"\n\nError Log File Saved to:\n{stats['error_log_path']}"

        messagebox.showinfo("Done", msg)


    def _time_leave_error(self, err):
        self.btn_tl.configure(state="normal")
        self.tl_prog.set(0)
        self.tl_prog_lbl.configure(text=f"Error: {err}")
        ui.update_status(self.tl_dot, self.tl_lbl, "Failed", "error")
        messagebox.showerror("Error", err)

    def _run_absent_job(self):
        try:
            import absent_management
            out_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile="Absent_Report.xlsx")
            if not out_path:
                self.after(0, lambda: self._absent_error("No output path selected"))
                return
            absent_management.process_absent_management(self.absent_emp_file.get(), self.absent_att_file.get(), self.absent_wfh_file.get(), out_path)
            self.after(0, lambda: self._absent_success(out_path))
        except Exception as e:
            self.after(0, lambda: self._absent_error(str(e)))

    def _run_att_job(self):
        try:
            import attendance_summary
            out_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile="Attendance_Summary.xlsx")
            if not out_path:
                self.after(0, lambda: self._att_error("No output path selected"))
                return
            attendance_summary.process_attendance_summary(self.att_summary_file.get(), out_path)
            self.after(0, lambda: self._att_success(out_path))
        except Exception as e:
            self.after(0, lambda: self._att_error(str(e)))

    def _run_generate_job(self):
        try:
            mode = self.gen_mode_seg.get() if hasattr(self, 'gen_mode_seg') else "Existing KRA Upload"
            
            if mode == "OKR Upload":
                from generate_upload import generate_okr_upload_file
                input_path = Path(self.selected_okr_input_file.get())
            else:
                from generate_upload import generate_upload_file
                input_path = Path(self.selected_gen_input_file.get())
            
            downloads = Path.home() / "Downloads"
            downloads.mkdir(parents=True, exist_ok=True)
            output_path = downloads / input_path.name

            mapping_text = self.gen_mapping_text.get("0.0", "end").strip()
            if not mapping_text:
                raise ValueError("Employee mapping cannot be empty.")
            
            employees = {}
            
            for line in mapping_text.splitlines():
                line = line.strip()
                if not line:
                    continue
                parts = line.split(",")
                if len(parts) == 3:
                    emp_id = parts[0].strip()
                    emp_name = parts[1].strip()
                    sheet_name = parts[2].strip()
                    
                    if sheet_name not in employees:
                        employees[sheet_name] = []
                    employees[sheet_name].append((emp_id, emp_name))
                else:
                    raise ValueError(f"Invalid format: '{line}'. Expected 'EMP ID, Name, Designation'.")

            if mode == "OKR Upload":
                generate_okr_upload_file(str(input_path), str(output_path), employees)
            else:
                generate_upload_file(str(input_path), str(output_path), employees)

            self.after(0, lambda: self._generate_success(output_path))

        except Exception as exc:
            import traceback
            error_text = f"{exc}\n\n{traceback.format_exc()}"
            self.after(0, lambda: self._generate_error(error_text))

    def _run_transform_job(self):
        try:
            input_path = Path(self.selected_input_file.get())
            output_path = self._build_output_path(input_path)

            transformed_sheets = {}

            with pd.ExcelFile(input_path) as excel_file:
                for sheet_name in excel_file.sheet_names:
                    df = pd.read_excel(excel_file, sheet_name=sheet_name, header=None)

                    transformed_df = App.transform_kra_dataframe(df)
                    if transformed_df is not None and not transformed_df.empty:
                        transformed_sheets[sheet_name] = transformed_df

            if not transformed_sheets:
                raise ValueError("No valid non-empty sheets found to transform.")

            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                for sheet_name, t_df in transformed_sheets.items():
                    safe_name = str(sheet_name)[:31] if sheet_name else "Sheet1"
                    export_df = t_df[["KRA", "KPI", "Weightage", "KPI %", "KRA %"]]
                    export_df.to_excel(writer, sheet_name=safe_name, index=False)

            self.apply_excel_formatting(output_path)
            self.after(0, lambda: self._transform_success(output_path, len(transformed_sheets)))

        except Exception as exc:
            error_text = f"{exc}\n\n{traceback.format_exc()}"
            self.after(0, lambda: self._transform_error(error_text))

    def _build_output_path(self, input_path: Path) -> Path:
        downloads = Path.home() / "Downloads"
        downloads.mkdir(parents=True, exist_ok=True)
        return downloads / input_path.name

    @staticmethod
    def is_effectively_empty(df: pd.DataFrame) -> bool:
        if df is None or df.empty:
            return True
        temp = df.copy().dropna(how="all")
        return temp.empty

    @staticmethod
    def is_serial_number_series(series: pd.Series) -> bool:
        non_null = series.dropna()
        if non_null.empty:
            return False

        numeric = pd.to_numeric(non_null, errors="coerce").dropna()
        if numeric.empty:
            return False

        # Ensure that at least 50% of the column is numeric so we don't 
        # misclassify text columns that happen to have a few sequential numbers
        if len(numeric) < len(non_null) * 0.5:
            return False

        values = numeric.astype(int).tolist()
        if len(values) < 2:
            return values[0] in [0, 1]

        sequential_matches = 0
        for i in range(1, len(values)):
            if values[i] == values[i - 1] + 1:
                sequential_matches += 1

        return sequential_matches >= max(1, len(values) - 2)

    @staticmethod
    def first_non_null(series: pd.Series):
        non_null = series.dropna()
        return non_null.iloc[0] if not non_null.empty else None

    @staticmethod
    def normalize_kra_pct(value):
        if pd.isna(value) or value == "":
            return ""
        try:
            num = float(value)
            if num <= 1.0:
                return int(round(num * 100))
            return int(round(num))
        except Exception:
            return ""

    @classmethod
    def transform_kra_dataframe(cls, df: pd.DataFrame):
        if cls.is_effectively_empty(df):
            return None

        df = df.copy()
        df = df.dropna(how="all").reset_index(drop=True)

        if df.shape[1] < 3:
            return None

        header_row_idx = -1
        kra_col_idx = -1
        kpi_col_idx = -1
        w_cols = []

        # 1. Dynamically find header row and columns based on keywords
        for i in range(min(30, len(df))):
            row_vals = df.iloc[i].astype(str).str.lower().str.strip()
            
            kc, pc = -1, -1
            wc = []
            
            for col_idx, val in enumerate(row_vals):
                if pd.isna(val) or len(val) < 2:
                    continue
                
                # Identify columns
                if kc == -1 and any(k in val for k in ["kra", "key result", "objective", "goal", "responsibility"]):
                    kc = col_idx
                elif pc == -1 and any(k in val for k in ["kpi", "key performance", "indicator", "measure", "metric", "parameter"]):
                    pc = col_idx
                elif any(k in val for k in ["weight", "wt", "wgt", "%", "proportion"]):
                    wc.append(col_idx)
            
            if kc != -1 and pc != -1:
                header_row_idx = i
                kra_col_idx = kc
                kpi_col_idx = pc
                w_cols = wc
                break

        # 2. Extract the appropriate columns
        if header_row_idx != -1:
            data_df = df.iloc[header_row_idx + 1:].reset_index(drop=True)
            kra_col = data_df.iloc[:, kra_col_idx]
            kpi_col = data_df.iloc[:, kpi_col_idx]
            
            if len(w_cols) >= 2:
                col3 = data_df.iloc[:, w_cols[0]]
                col4 = data_df.iloc[:, w_cols[1]]
            elif len(w_cols) == 1:
                col3 = data_df.iloc[:, w_cols[0]]
                col4 = None
            else:
                # Fallback to columns immediately following KPI if no weight keywords found
                next_cols = [c for c in range(data_df.shape[1]) if c > kpi_col_idx and c not in [kra_col_idx, kpi_col_idx]]
                if len(next_cols) >= 2:
                    col3 = data_df.iloc[:, next_cols[0]]
                    col4 = data_df.iloc[:, next_cols[1]]
                elif len(next_cols) == 1:
                    col3 = data_df.iloc[:, next_cols[0]]
                    col4 = None
                else:
                    return None
        else:
            # 3. Fallback to positional logic if headers are not explicitly found
            start_idx = 0
            if cls.is_serial_number_series(df.iloc[:, 0]):
                start_idx = 1
                
            if df.shape[1] - start_idx < 3:
                return None
                
            # Assume row 0 is the header (since we read with header=None) and skip it
            data_df = df.iloc[1:].reset_index(drop=True)
            
            kra_col = data_df.iloc[:, start_idx]
            kpi_col = data_df.iloc[:, start_idx + 1]
            col3 = data_df.iloc[:, start_idx + 2]
            col4 = data_df.iloc[:, start_idx + 3] if data_df.shape[1] - start_idx >= 4 else None

        # CASE A: 4 or more usable columns
        # KRA | KPI | KRA Weightage | KPI Weightage
        if col4 is not None:
            working = pd.DataFrame({
                "KRA_RAW": kra_col,
                "KPI_RAW": kpi_col,
                "KRA_W_RAW": col3,
                "KPI_W_RAW": col4,
            })

            working["KRA"] = working["KRA_RAW"].ffill()
            working["KRA"] = working["KRA"].astype(str).str.strip()
            working["KPI"] = working["KPI_RAW"].astype(str).str.strip()

            working = working[
                (working["KRA"].notna()) &
                (working["KRA"] != "") &
                (working["KRA"].str.lower() != "nan") &
                (working["KPI"].notna()) &
                (working["KPI"] != "") &
                (working["KPI"].str.lower() != "nan")
            ].copy()

            if working.empty:
                return None

            output_rows = []

            for kra_name, group in working.groupby("KRA", sort=False):
                group = group.copy().reset_index(drop=True)

                raw_kra_pct_value = cls.first_non_null(group["KRA_W_RAW"])
                kra_pct = cls.normalize_kra_pct(raw_kra_pct_value)

                numeric_weights = pd.to_numeric(group["KPI_W_RAW"], errors="coerce").fillna(0)
                total_w = numeric_weights.sum()

                if total_w > 0:
                    kpi_pcts = [round((w / total_w) * 100, 2) for w in numeric_weights.tolist()]
                    
                    diff = round(100.0 - sum(kpi_pcts), 2)
                    if kpi_pcts:
                        kpi_pcts[-1] = round(kpi_pcts[-1] + diff, 2)
                else:
                    kpi_pcts = [0.0] * len(group)

                for i, (_, row) in enumerate(group.iterrows()):
                    raw_weight = row["KPI_W_RAW"]
                    if pd.isna(raw_weight):
                        raw_weight = ""

                    output_rows.append({
                        "KRA": kra_name,
                        "KPI": row["KPI"],
                        "Weightage": raw_weight,
                        "KPI %": round(kpi_pcts[i], 2),
                        "KRA %": kra_pct,
                    })

            return pd.DataFrame(output_rows)

        # CASE B: exactly 3 usable columns
        # KRA | KPI | Row Weightage
        working = pd.DataFrame({
            "KRA_RAW": kra_col,
            "KPI_RAW": kpi_col,
            "ROW_W_RAW": col3,
        })

        working["KRA"] = working["KRA_RAW"].ffill()
        working["KRA"] = working["KRA"].astype(str).str.strip()
        working["KPI"] = working["KPI_RAW"].astype(str).str.strip()

        working = working[
            (working["KRA"].notna()) &
            (working["KRA"] != "") &
            (working["KRA"].str.lower() != "nan") &
            (working["KPI"].notna()) &
            (working["KPI"] != "") &
            (working["KPI"].str.lower() != "nan")
        ].copy()

        if working.empty:
            return None

        output_rows = []

        for kra_name, group in working.groupby("KRA", sort=False):
            group = group.copy().reset_index(drop=True)

            numeric_weights = pd.to_numeric(group["ROW_W_RAW"], errors="coerce").fillna(0)
            total_w = numeric_weights.sum()

            if total_w > 0:
                kpi_pcts = [round((w / total_w) * 100, 2) for w in numeric_weights.tolist()]
                
                diff = round(100.0 - sum(kpi_pcts), 2)
                if kpi_pcts:
                    kpi_pcts[-1] = round(kpi_pcts[-1] + diff, 2)
            else:
                kpi_pcts = [0.0] * len(group)

            kra_pct = cls.normalize_kra_pct(total_w)

            for i, (_, row) in enumerate(group.iterrows()):
                raw_weight = row["ROW_W_RAW"]
                if pd.isna(raw_weight):
                    raw_weight = ""

                output_rows.append({
                    "KRA": kra_name,
                    "KPI": row["KPI"],
                    "Weightage": raw_weight,
                    "KPI %": round(kpi_pcts[i], 2),
                    "KRA %": kra_pct,
                })

        return pd.DataFrame(output_rows)

    @staticmethod
    def apply_excel_formatting(output_path: Path):
        wb = load_workbook(output_path)

        header_fill = PatternFill(fill_type="solid", start_color="1F4E79", end_color="1F4E79")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        white_fill = PatternFill(fill_type="solid", start_color="FFFFFF", end_color="FFFFFF")
        blue_fill = PatternFill(fill_type="solid", start_color="DCE6F1", end_color="DCE6F1")
        yellow_fill = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")

        left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        row_font = Font(name="Calibri", size=10)

        thin_side = Side(style="thin", color="BFBFBF")
        thin_border = Border() # No border as requested

        for ws in wb.worksheets:
            headers = ["KRA", "KPI", "Weightage", "KPI %", "KRA %"]
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = thin_border

            ws.row_dimensions[1].height = 30

            ws.column_dimensions["A"].width = 50
            ws.column_dimensions["B"].width = 80
            ws.column_dimensions["C"].width = 12
            ws.column_dimensions["D"].width = 10
            ws.column_dimensions["E"].width = 10

            current_fill_name = "white"
            previous_kra = None

            max_row = ws.max_row

            for row_num in range(2, max_row + 1):
                kra_value = ws.cell(row=row_num, column=1).value
                kra_pct_value = ws.cell(row=row_num, column=5).value

                if kra_value != previous_kra:
                    if previous_kra is not None:
                        current_fill_name = "blue" if current_fill_name == "white" else "white"
                    previous_kra = kra_value

                row_fill = white_fill if current_fill_name == "white" else blue_fill

                if kra_pct_value in [None, ""]:
                    row_fill = yellow_fill

                for col_num in range(1, 6):
                    cell = ws.cell(row=row_num, column=col_num)
                    cell.fill = row_fill
                    cell.border = thin_border
                    cell.font = row_font

                    if col_num in [1, 2]:
                        cell.alignment = left_alignment
                    else:
                        cell.alignment = center_alignment

                    if col_num == 4:
                        cell.number_format = "0.00"

                ws.row_dimensions[row_num].height = 30

        wb.save(output_path)


if __name__ == "__main__":
    app = App()
    app.mainloop()
